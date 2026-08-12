"""Deterministically render validated CV content into a DOCX template."""

from __future__ import annotations

import re
import unicodedata
import zipfile
from collections.abc import Collection
from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import urlsplit
from xml.etree import ElementTree

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.text.paragraph import Paragraph

from app.ai.schema_models import CvContent, CvSectionKind
from app.artefacts import MAX_SEGMENT_LENGTH, ArtefactStore
from app.settings import CvLayoutSettings
from app.template_validation import (
    TemplateStyleError,
    validate_template_styles,
)

WINDOWS_RESERVED_NAMES = {
    "aux",
    "clock$",
    "con",
    "nul",
    "prn",
    *(f"com{number}" for number in range(1, 10)),
    *(f"lpt{number}" for number in range(1, 10)),
}
RELATIONSHIP_NAMESPACE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
)
RELATIONSHIP_REFERENCE_ATTRIBUTES = frozenset(
    {
        f"{{{RELATIONSHIP_NAMESPACE}}}id",
        f"{{{RELATIONSHIP_NAMESPACE}}}embed",
        f"{{{RELATIONSHIP_NAMESPACE}}}link",
    }
)
CONTENT_TYPES_NAMESPACE = (
    "http://schemas.openxmlformats.org/package/2006/content-types"
)
MARKUP_COMPATIBILITY_NAMESPACE = (
    "http://schemas.openxmlformats.org/markup-compatibility/2006"
)
XML_CONTENT_TYPES = frozenset(
    {
        "application/vnd.ms-office.vmldrawing",
        "application/vnd.openxmlformats-officedocument.vmldrawing",
        "application/xml",
        "text/xml",
    }
)
PACKAGE_ROOT_RELATIONSHIP_KINDS = frozenset(
    {
        "core-properties",
        "officeDocument",
    }
)
RENDERING_RELATIONSHIP_KINDS = frozenset(
    {
        "fontTable",
        "numbering",
        "settings",
        "styles",
        "stylesWithEffects",
        "theme",
        "webSettings",
    }
)
OPTIONAL_CONTENT_RELATIONSHIP_KINDS = frozenset(
    {
        "comments",
        "commentsExtended",
        "custom-properties",
        "endnotes",
        "extended-properties",
        "footnotes",
        "glossaryDocument",
        "people",
        "thumbnail",
    }
)
CORE_PROPERTY_TEXT_LIMIT = 255
DETERMINISTIC_ZIP_TIMESTAMP = (2000, 1, 1, 0, 0, 0)
MAX_CANDIDATE_BYTES = 234
STORY_PART_PATTERN = re.compile(r"word/(?:document|header\d+|footer\d+)\.xml$")


class WordWriterError(ValueError):
    """Raised when a DOCX cannot be rendered safely and deterministically."""


@dataclass(frozen=True)
class CandidateDocument:
    """The persisted DOCX candidate and its content hash."""

    filename: str
    relative_path: str
    sha256: str


@dataclass(frozen=True)
class _PackageContentTypes:
    defaults: dict[str, str]
    overrides: dict[str, str]

    @classmethod
    def from_xml(cls, data: bytes) -> _PackageContentTypes:
        root = ElementTree.fromstring(data)
        default_tag = f"{{{CONTENT_TYPES_NAMESPACE}}}Default"
        override_tag = f"{{{CONTENT_TYPES_NAMESPACE}}}Override"
        defaults = {
            element.get("Extension", "").casefold(): element.get(
                "ContentType", ""
            )
            for element in root.findall(default_tag)
        }
        overrides = {
            element.get("PartName", "").lstrip("/"): element.get(
                "ContentType", ""
            )
            for element in root.findall(override_tag)
        }
        return cls(defaults=defaults, overrides=overrides)

    def content_type(self, part_name: str) -> str | None:
        override = self.overrides.get(part_name)
        if override is not None:
            return override
        filename = part_name.rpartition("/")[2]
        _, separator, extension = filename.rpartition(".")
        if not separator:
            return None
        return self.defaults.get(extension.casefold())

    def is_xml_backed(self, part_name: str) -> bool:
        if part_name.casefold().endswith(".xml"):
            return True
        content_type = self.content_type(part_name)
        if content_type is None:
            return False
        normalized = content_type.partition(";")[0].strip().casefold()
        return normalized.endswith("+xml") or normalized in XML_CONTENT_TYPES


def _safe_filename_part(value: str, fallback: str) -> str:
    normalized = unicodedata.normalize("NFKC", value).strip()
    cleaned = "".join(
        character
        if character.isalnum() or character in {"-", "_", ".", " ", "(", ")"}
        else " "
        for character in normalized
    )
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" .")
    if not cleaned:
        cleaned = fallback
    if cleaned.split(".", maxsplit=1)[0].casefold() in WINDOWS_RESERVED_NAMES:
        cleaned = f"_{cleaned}"
    return cleaned


def _truncate_utf8(
    value: str, maximum_characters: int, maximum_bytes: int
) -> str:
    """Truncate text without splitting a Unicode code point."""
    if maximum_characters <= 0 or maximum_bytes <= 0:
        return ""
    candidate = value[:maximum_characters]
    encoded = candidate.encode("utf-8")
    if len(encoded) <= maximum_bytes:
        return candidate
    return encoded[:maximum_bytes].decode("utf-8", errors="ignore")


def candidate_filename(full_name: str, role: str) -> str:
    """Build the required portable candidate filename."""
    name = _safe_filename_part(full_name, "Candidate")
    target_role = _safe_filename_part(role, "Role")
    suffix = ".docx"
    separator = " - "
    maximum_characters = MAX_SEGMENT_LENGTH - len(separator) - len(suffix)
    maximum_bytes = (
        MAX_CANDIDATE_BYTES
        - len(separator.encode("utf-8"))
        - len(suffix.encode("utf-8"))
    )
    reserved_role_character = target_role[0]
    name = (
        _truncate_utf8(
            name,
            maximum_characters - 1,
            maximum_bytes - len(reserved_role_character.encode("utf-8")),
        ).rstrip(" .")
        or "C"
    )
    role_characters = maximum_characters - len(name)
    role_bytes = maximum_bytes - len(name.encode("utf-8"))
    target_role = (
        _truncate_utf8(target_role, role_characters, role_bytes).rstrip(" .")
        or "R"
    )
    return f"{name}{separator}{target_role}{suffix}"


def _set_style_font(style: Any, name: str, size: float) -> None:
    font = style.font
    font.name = name
    font.size = Pt(size)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), name)


def _validate_styles(document: DocumentObject) -> None:
    try:
        validate_template_styles(document)
    except TemplateStyleError as error:
        raise WordWriterError(str(error)) from error


def _clear_story(story: Any) -> None:
    story_element = story._element
    for child in list(story_element):
        story_element.remove(child)
    story.add_paragraph()


def _clear_template(document: DocumentObject) -> None:
    document_element = document._element
    for child in list(document_element):
        if child.tag == qn("w:background"):
            document_element.remove(child)
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    story_names = (
        "header",
        "first_page_header",
        "even_page_header",
        "footer",
        "first_page_footer",
        "even_page_footer",
    )
    cleared_parts: set[str] = set()
    for section in document.sections:
        for story_name in story_names:
            story = getattr(section, story_name)
            part_name = str(story.part.partname)
            if part_name not in cleared_parts:
                _clear_story(story)
                cleared_parts.add(part_name)
    for section in document.sections:
        section_properties = section._sectPr
        section_properties.attrib.clear()
        for child in tuple(section_properties):
            section_properties.remove(child)


def _apply_page_layout(
    document: DocumentObject, layout: CvLayoutSettings
) -> None:
    for section in document.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        if layout.page.size == "A4":
            section.page_width = Mm(210)
            section.page_height = Mm(297)
        else:
            section.page_width = Inches(8.5)
            section.page_height = Inches(11)
        section.top_margin = Inches(layout.page.margins.top)
        section.right_margin = Inches(layout.page.margins.right)
        section.bottom_margin = Inches(layout.page.margins.bottom)
        section.left_margin = Inches(layout.page.margins.left)


def _source_part_name(relationship_name: str) -> str | None:
    if relationship_name == "_rels/.rels":
        return None
    if relationship_name.startswith("_rels/"):
        filename = relationship_name.removeprefix("_rels/")
        if filename.endswith(".rels"):
            return filename.removesuffix(".rels")
    parent, _, filename = relationship_name.rpartition("/_rels/")
    if not parent or not filename.endswith(".rels"):
        return None
    return f"{parent}/{filename.removesuffix('.rels')}"


def _relationship_name(source_part: str | None) -> str:
    if source_part is None:
        return "_rels/.rels"
    parent, separator, filename = source_part.rpartition("/")
    if not separator:
        return f"_rels/{filename}.rels"
    return f"{parent}/_rels/{filename}.rels"


def _relationship_target(source: str, target: str) -> str:
    normalized_target = target.replace("\\", "/")
    target_parts = normalized_target.split("/")
    source_parts = (
        [] if normalized_target.startswith("/") else source.split("/")[:-1]
    )
    for part in target_parts:
        if not part or part == ".":
            continue
        if part == "..":
            if source_parts:
                source_parts.pop()
            continue
        source_parts.append(part)
    return "/".join(source_parts)


def _relationship_kind(relationship_type: str) -> str:
    return relationship_type.rpartition("/")[2]


def _retain_unreferenced_relationship(
    source_part: str | None, relationship_type: str
) -> bool:
    relationship_kind = _relationship_kind(relationship_type)
    if relationship_kind in OPTIONAL_CONTENT_RELATIONSHIP_KINDS:
        return False
    if source_part is None:
        return relationship_kind in PACKAGE_ROOT_RELATIONSHIP_KINDS
    return relationship_kind in RENDERING_RELATIONSHIP_KINDS


def _referenced_relationship_ids(data: bytes) -> set[str]:
    root = ElementTree.fromstring(data)
    return {
        value
        for element in root.iter()
        for attribute, value in element.attrib.items()
        if attribute in RELATIONSHIP_REFERENCE_ATTRIBUTES
    }


def _sanitize_settings(document: DocumentObject) -> None:
    settings = document.settings.element
    private_tags = {
        qn("w:attachedTemplate"),
        qn("w:docVars"),
        qn("w:mailMerge"),
    }
    for parent in settings.iter():
        for child in tuple(parent):
            if child.tag in private_tags:
                parent.remove(child)


def _element_values(
    root: ElementTree.Element,
    element_name: str,
    attribute_name: str,
) -> set[str]:
    return {
        value
        for element in root.iter(qn(element_name))
        if (value := element.get(qn(attribute_name))) is not None
    }


def _numbering_children(
    parent: ElementTree.Element, element_name: str
) -> list[ElementTree.Element]:
    compatibility_tags = {
        f"{{{MARKUP_COMPATIBILITY_NAMESPACE}}}Choice",
        f"{{{MARKUP_COMPATIBILITY_NAMESPACE}}}Fallback",
    }
    alternate_content_tag = (
        f"{{{MARKUP_COMPATIBILITY_NAMESPACE}}}AlternateContent"
    )
    children: list[ElementTree.Element] = []
    for child in parent:
        if child.tag == qn(element_name):
            children.append(child)
        elif child.tag == alternate_content_tag:
            for branch in child:
                if branch.tag in compatibility_tags:
                    children.extend(_numbering_children(branch, element_name))
    return children


def _prune_numbering_resources(members: dict[str, bytes]) -> None:
    styles_name = "word/styles.xml"
    numbering_name = "word/numbering.xml"
    if styles_name not in members or numbering_name not in members:
        return

    styles = ElementTree.fromstring(members[styles_name])
    numbering = ElementTree.fromstring(members[numbering_name])
    style_definitions = {
        style_id: style
        for style in styles.iter(qn("w:style"))
        if (style_id := style.get(qn("w:styleId"))) is not None
    }
    abstract_definitions = {
        abstract_id: abstract_number
        for abstract_number in numbering.findall(qn("w:abstractNum"))
        if (abstract_id := abstract_number.get(qn("w:abstractNumId")))
        is not None
    }
    number_definitions = {
        num_id: number
        for number in numbering.findall(qn("w:num"))
        if (num_id := number.get(qn("w:numId"))) is not None
    }
    live_style_ids: set[str] = set()
    live_num_ids: set[str] = set()
    for name, data in members.items():
        if not STORY_PART_PATTERN.fullmatch(name):
            continue
        story = ElementTree.fromstring(data)
        live_style_ids.update(_element_values(story, "w:pStyle", "w:val"))
        live_num_ids.update(_element_values(story, "w:numId", "w:val"))

    visited_style_ids: set[str] = set()
    live_abstract_ids: set[str] = set()
    while True:
        pending_styles = live_style_ids - visited_style_ids
        for style_id in pending_styles:
            visited_style_ids.add(style_id)
            style = style_definitions.get(style_id)
            if style is None:
                continue
            live_style_ids.update(_element_values(style, "w:basedOn", "w:val"))
            live_num_ids.update(_element_values(style, "w:numId", "w:val"))

        previous_abstract_ids = set(live_abstract_ids)
        for num_id in live_num_ids:
            number = number_definitions.get(num_id)
            if number is None:
                continue
            abstract_reference = number.find(qn("w:abstractNumId"))
            if abstract_reference is None:
                continue
            abstract_id = abstract_reference.get(qn("w:val"))
            if abstract_id is not None:
                live_abstract_ids.add(abstract_id)
        for abstract_id in live_abstract_ids:
            abstract_number = abstract_definitions.get(abstract_id)
            if abstract_number is None:
                continue
            live_style_ids.update(
                _element_values(abstract_number, "w:numStyleLink", "w:val")
            )
            live_style_ids.update(
                _element_values(abstract_number, "w:styleLink", "w:val")
            )
        if (
            not (live_style_ids - visited_style_ids)
            and live_abstract_ids == previous_abstract_ids
        ):
            break

    live_picture_bullet_ids: set[str] = set()
    for num_id in live_num_ids:
        number = number_definitions.get(num_id)
        if number is None:
            continue
        abstract_reference = number.find(qn("w:abstractNumId"))
        if abstract_reference is None:
            continue
        abstract_id = abstract_reference.get(qn("w:val"))
        if abstract_id is None:
            continue
        abstract_number = abstract_definitions.get(abstract_id)
        if abstract_number is None:
            continue
        abstract_levels: dict[str, list[ElementTree.Element]] = {}
        for level in _numbering_children(abstract_number, "w:lvl"):
            level_id = level.get(qn("w:ilvl"))
            if level_id is not None:
                abstract_levels.setdefault(level_id, []).append(level)
        level_overrides: dict[str, list[ElementTree.Element]] = {}
        for level_override in _numbering_children(number, "w:lvlOverride"):
            level_id = level_override.get(qn("w:ilvl"))
            if level_id is not None:
                level_overrides.setdefault(level_id, []).append(level_override)
        for level_id in abstract_levels.keys() | level_overrides.keys():
            override_levels = [
                level
                for level_override in level_overrides.get(level_id, [])
                for level in _numbering_children(level_override, "w:lvl")
            ]
            effective_levels = override_levels or abstract_levels.get(
                level_id, []
            )
            for level in effective_levels:
                live_picture_bullet_ids.update(
                    _element_values(level, "w:lvlPicBulletId", "w:val")
                )

    for child in tuple(numbering):
        if child.tag == qn("w:num"):
            if child.get(qn("w:numId")) not in live_num_ids:
                numbering.remove(child)
        elif child.tag == qn("w:abstractNum"):
            if child.get(qn("w:abstractNumId")) not in live_abstract_ids:
                numbering.remove(child)
        elif child.tag == qn("w:numPicBullet") and (
            child.get(qn("w:numPicBulletId")) not in live_picture_bullet_ids
        ):
            numbering.remove(child)
    members[numbering_name] = ElementTree.tostring(
        numbering,
        encoding="utf-8",
        xml_declaration=True,
    )


def _internal_relationship_targets(
    members: dict[str, bytes], source_part: str | None
) -> set[str]:
    relationship_name = _relationship_name(source_part)
    if relationship_name not in members:
        return set()
    relationships = ElementTree.fromstring(members[relationship_name])
    source = source_part or ""
    return {
        _relationship_target(source, relationship.get("Target", ""))
        for relationship in relationships
        if relationship.get("TargetMode") != "External"
    }


def _sanitize_package(data: bytes) -> bytes:
    source = BytesIO(data)
    destination = BytesIO()
    with zipfile.ZipFile(source) as archive:
        members = {
            member.filename: archive.read(member.filename)
            for member in archive.infolist()
        }

    package_content_types = _PackageContentTypes.from_xml(
        members["[Content_Types].xml"]
    )
    _prune_numbering_resources(members)
    referenced_ids: dict[str, set[str]] = {}
    for name, member_data in members.items():
        if name.endswith(".rels") or not package_content_types.is_xml_backed(
            name
        ):
            continue
        try:
            referenced_ids[name] = _referenced_relationship_ids(member_data)
        except ElementTree.ParseError:
            referenced_ids[name] = set()
    for name, relationship_data in tuple(members.items()):
        if not name.endswith(".rels"):
            continue
        source_part = _source_part_name(name)
        relationships = ElementTree.fromstring(relationship_data)
        used_ids = (
            referenced_ids.get(source_part, set())
            if source_part is not None
            else set()
        )
        for relationship in tuple(relationships):
            relationship_type = relationship.get("Type", "")
            relationship_id = relationship.get("Id", "")
            if (
                relationship_id not in used_ids
                and not _retain_unreferenced_relationship(
                    source_part, relationship_type
                )
            ):
                relationships.remove(relationship)
        members[name] = ElementTree.tostring(
            relationships,
            encoding="utf-8",
            xml_declaration=True,
        )

    reachable_parts: set[str] = set()
    pending_parts = list(_internal_relationship_targets(members, None))
    while pending_parts:
        part_name = pending_parts.pop()
        if part_name in reachable_parts or part_name not in members:
            continue
        reachable_parts.add(part_name)
        pending_parts.extend(_internal_relationship_targets(members, part_name))

    removed_parts = {
        name
        for name in members
        if name != "[Content_Types].xml"
        and not name.endswith(".rels")
        and name not in reachable_parts
    }
    for name in tuple(members):
        if name in removed_parts:
            del members[name]
            continue
        source_part = _source_part_name(name)
        if (
            name.endswith(".rels")
            and source_part is not None
            and (source_part not in reachable_parts)
        ):
            del members[name]

    content_types_name = "[Content_Types].xml"
    content_types_root = ElementTree.fromstring(members[content_types_name])
    for override in tuple(content_types_root):
        part_name = override.get("PartName", "").lstrip("/")
        if part_name in removed_parts:
            content_types_root.remove(override)
    members[content_types_name] = ElementTree.tostring(
        content_types_root,
        encoding="utf-8",
        xml_declaration=True,
    )

    with zipfile.ZipFile(
        destination,
        "w",
        zipfile.ZIP_DEFLATED,
        compresslevel=9,
    ) as archive:
        for name in sorted(members):
            archive.writestr(_deterministic_zip_info(name), members[name])
    return destination.getvalue()


def _deterministic_zip_info(name: str) -> zipfile.ZipInfo:
    member = zipfile.ZipInfo(name, date_time=DETERMINISTIC_ZIP_TIMESTAMP)
    member.compress_type = zipfile.ZIP_DEFLATED
    member.create_system = 3
    member.external_attr = 0o600 << 16
    member.internal_attr = 0
    member.extra = b""
    member.comment = b""
    return member


def _apply_styles(document: DocumentObject, layout: CvLayoutSettings) -> None:
    normal = document.styles["Normal"]
    title = document.styles["Title"]
    heading = document.styles["Heading 1"]
    bullet = document.styles["List Bullet"]
    _set_style_font(normal, layout.fonts.body, layout.fonts.body_size_pt)
    _set_style_font(bullet, layout.fonts.body, layout.fonts.body_size_pt)
    _set_style_font(
        heading, layout.fonts.headings, layout.fonts.heading_size_pt
    )
    _set_style_font(
        title, layout.fonts.headings, layout.fonts.heading_size_pt + 4
    )
    heading.font.color.rgb = RGBColor.from_string(
        layout.styles.heading_color.upper()
    )
    title.font.color.rgb = RGBColor.from_string(
        layout.styles.accent_color.upper()
    )
    for style in (normal, bullet):
        paragraph_format = style.paragraph_format
        paragraph_format.line_spacing = layout.spacing.line
        paragraph_format.space_after = Pt(layout.spacing.paragraph_after_pt)
    bullet.paragraph_format.left_indent = Inches(
        layout.styles.bullet_indent_inches
    )
    heading.paragraph_format.space_after = Pt(layout.spacing.section_after_pt)
    heading.paragraph_format.keep_with_next = True


def _hyperlink_target(text: str, kind: str) -> str | None:
    if kind == "email":
        return f"mailto:{text}" if "@" in text else None
    if kind == "phone":
        compact = re.sub(r"[^0-9+]", "", text)
        return f"tel:{compact}" if compact else None
    if kind == "website":
        parsed = urlsplit(text)
        if parsed.scheme.lower() in {"http", "https"} and parsed.netloc:
            return text
    return None


def _add_hyperlink(
    paragraph: Paragraph, text: str, target: str, color_value: str
) -> None:
    relationship_id = paragraph.part.relate_to(
        target,
        "http://schemas.openxmlformats.org/officeDocument/2006/"
        "relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), color_value.upper())
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    run.append(properties)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_contact_line(
    document: DocumentObject,
    content: CvContent,
    layout: CvLayoutSettings,
) -> None:
    identity = content.identity
    values = (
        ("professional_title", identity.professional_title),
        ("email", identity.email),
        ("phone", identity.phone),
        ("location", identity.location),
        ("website", identity.website),
    )
    present = tuple(
        (kind, value) for kind, value in values if value is not None
    )
    if not present:
        return
    paragraph = document.add_paragraph(style="Normal")
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    for index, (kind, value) in enumerate(present):
        if index:
            paragraph.add_run(" | ")
        target = _hyperlink_target(value.text, kind)
        if target is None:
            paragraph.add_run(value.text)
        else:
            _add_hyperlink(
                paragraph,
                value.text,
                target,
                layout.styles.accent_color,
            )


def _add_sections(
    document: DocumentObject,
    content: CvContent,
    layout: CvLayoutSettings,
) -> None:
    order = {kind: index for index, kind in enumerate(layout.section_order)}
    sections = sorted(
        content.sections, key=lambda section: order[section.kind.value]
    )
    for section in sections:
        document.add_paragraph(section.heading, style="Heading 1")
        claim_style = (
            "Normal" if section.kind is CvSectionKind.SUMMARY else "List Bullet"
        )
        for claim in section.claims:
            document.add_paragraph(claim.text, style=claim_style)


def _safe_core_property_text(value: str) -> str:
    safe = "".join(
        character
        for character in value
        if character in "\t\n\r"
        or "\x20" <= character <= "\ud7ff"
        or "\ue000" <= character <= "\ufffd"
        or "\U00010000" <= character <= "\U0010ffff"
    )
    return safe[:CORE_PROPERTY_TEXT_LIMIT]


def _metadata_title(full_name: str, role: str) -> str:
    safe_name = _safe_core_property_text(full_name)
    separator = " - "
    role_limit = CORE_PROPERTY_TEXT_LIMIT - len(safe_name) - len(separator)
    if role_limit <= 0:
        return safe_name[:CORE_PROPERTY_TEXT_LIMIT]
    safe_role = _safe_core_property_text(role)[:role_limit]
    return f"{safe_name}{separator}{safe_role}"


def _set_safe_metadata(
    document: DocumentObject, content: CvContent, role: str
) -> None:
    properties = document.core_properties
    properties_element = properties._element
    for child in tuple(properties_element):
        properties_element.remove(child)

    properties.title = _metadata_title(content.identity.full_name.text, role)
    properties.subject = "Tailored curriculum vitae"
    properties.author = _safe_core_property_text(
        content.identity.full_name.text
    )
    properties.last_modified_by = "JobHunter"
    properties.comments = "Generated by JobHunter from validated CV content."
    properties.category = "CV"
    properties.language = "en"
    properties.version = "1"
    properties.revision = 1
    fixed_time = datetime(2000, 1, 1, tzinfo=UTC)
    properties.created = fixed_time
    properties.modified = fixed_time
    properties.last_printed = fixed_time


class WordWriter:
    """Render CV content into an atomic private DOCX candidate."""

    def __init__(
        self,
        artefacts: ArtefactStore,
        *,
        available_fonts: Collection[str] | None = None,
    ) -> None:
        self._artefacts = artefacts
        self._available_fonts = (
            {font.casefold() for font in available_fonts}
            if available_fonts is not None
            else None
        )

    def _validate_fonts(self, layout: CvLayoutSettings) -> None:
        if self._available_fonts is None:
            return
        missing = sorted(
            font
            for font in (layout.fonts.body, layout.fonts.headings)
            if font.casefold() not in self._available_fonts
        )
        if missing:
            raise WordWriterError(
                "Configured CV fonts are unavailable: " + ", ".join(missing)
            )

    def write(
        self,
        content: CvContent,
        role: str,
        base_path: str | Path,
        template_path: str | Path,
        layout: CvLayoutSettings,
    ) -> CandidateDocument:
        """Write one candidate atomically beneath the generation directory."""
        if not role or role != role.strip():
            raise WordWriterError("Application role must be trimmed text.")
        self._validate_fonts(layout)
        try:
            document = Document(str(template_path))
        except Exception as error:
            raise WordWriterError(
                "CV template is corrupt or unreadable."
            ) from error
        _validate_styles(document)
        _sanitize_settings(document)
        _clear_template(document)
        _apply_page_layout(document, layout)
        _apply_styles(document, layout)

        document.add_paragraph(
            content.identity.full_name.text, style="Title"
        ).alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        _add_contact_line(document, content, layout)
        _add_sections(document, content, layout)
        _set_safe_metadata(document, content, role)

        output = BytesIO()
        try:
            document.save(output)
        except Exception as error:
            raise WordWriterError(
                "CV candidate could not be rendered."
            ) from error
        filename = candidate_filename(content.identity.full_name.text, role)
        relative_path = (Path(base_path) / filename).as_posix()
        self._artefacts.resolve(relative_path)
        self._artefacts.write_bytes(
            relative_path, _sanitize_package(output.getvalue())
        )
        return CandidateDocument(
            filename=filename,
            relative_path=relative_path,
            sha256=self._artefacts.sha256(relative_path),
        )
