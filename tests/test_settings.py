from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
import yaml
from docx import Document
from docx.oxml.ns import qn

from app.settings import (
    MAX_PROFILE_BYTES,
    SettingsError,
    load_ai_settings,
    validate_layout_input,
    validate_private_inputs,
    validate_profile_input,
    validate_template_input,
)


def ai_values(**overrides: Any) -> dict[str, Any]:
    values: dict[str, Any] = {
        "provider": {
            "name": "openai",
            "assessment_model": "assessment-model",
            "cv_model": "cv-model",
            "reasoning_effort": "high",
            "request_timeout_seconds": 90,
        },
        "scoring": {"threshold": 75, "taxonomy_version": "v1"},
        "queue": {
            "poll_interval_seconds": 1,
            "heartbeat_interval_seconds": 10,
            "work_lease_seconds": 45,
            "resource_lease_seconds": 180,
            "concurrency": 3,
            "max_attempts": 2,
        },
    }
    values.update(overrides)
    return values


def layout_values(**overrides: Any) -> dict[str, Any]:
    values: dict[str, Any] = {
        "page": {
            "size": "A4",
            "orientation": "portrait",
            "margins": {
                "top": 0.5,
                "right": 0.5,
                "bottom": 0.5,
                "left": 0.5,
            },
        },
        "fonts": {
            "body": "Aptos",
            "headings": "Aptos Display",
            "body_size_pt": 9.5,
            "heading_size_pt": 12,
        },
        "spacing": {
            "line": 1.0,
            "paragraph_after_pt": 3,
            "section_after_pt": 6,
        },
        "styles": {
            "heading_color": "1F2937",
            "accent_color": "2563EB",
            "bullet_indent_inches": 0.2,
        },
        "output": {"filename": "Tailored CV.docx"},
    }
    values.update(overrides)
    return values


def write_yaml(path: Path, values: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(yaml.safe_dump(values), encoding="utf-8")


def write_private_inputs(root: Path) -> Path:
    profile_root = root / "private" / "profile"
    profile_root.mkdir(parents=True)
    (profile_root / "profile.md").write_text(
        "# Example Person\n\nPython engineer.\n", encoding="utf-8"
    )
    document = Document()
    document.add_paragraph("Template")
    document.save(str(profile_root / "cv-template.docx"))
    write_yaml(profile_root / "cv-layout.yaml", layout_values())
    return profile_root


def test_tracked_settings_load_defaults_and_typed_overrides(
    tmp_path: Path,
) -> None:
    config = tmp_path / "ai.yaml"
    write_yaml(config, ai_values())

    defaults = load_ai_settings(config)
    overridden = load_ai_settings(
        config,
        {
            "provider": {"assessment_model": "override-model"},
            "scoring": {"threshold": 82.5},
        },
    )

    assert defaults.provider.name == "openai"
    assert defaults.queue.concurrency == 3
    assert overridden.provider.assessment_model == "override-model"
    assert overridden.provider.cv_model == "cv-model"
    assert overridden.scoring.threshold == 82.5


def test_checked_in_settings_allow_slow_structured_generation() -> None:
    """Keep the production request window suitable for high reasoning."""
    settings = load_ai_settings()

    assert settings.provider.request_timeout_seconds == 300


@pytest.mark.parametrize(
    ("change", "message"),
    [
        ({"provider": {"name": "anthropic"}}, "AI settings are invalid"),
        (
            {
                "queue": {
                    "poll_interval_seconds": 10,
                    "heartbeat_interval_seconds": 10,
                }
            },
            "poll interval",
        ),
        (
            {
                "queue": {
                    "heartbeat_interval_seconds": 30,
                    "work_lease_seconds": 45,
                }
            },
            "work lease",
        ),
        (
            {"provider": {"reasoning_effort": "urgent"}},
            "AI settings are invalid",
        ),
        ({"queue": {"concurrency": 4}}, "AI settings are invalid"),
    ],
)
def test_tracked_settings_reject_unknown_provider_and_invalid_relationships(
    tmp_path: Path, change: dict[str, object], message: str
) -> None:
    config = tmp_path / "ai.yaml"
    write_yaml(config, ai_values())

    with pytest.raises(SettingsError, match=message):
        load_ai_settings(config, change)


@pytest.mark.parametrize(
    ("change", "message"),
    [
        ({"scoring": {"taxonomy_version": "v2"}}, "unsupported"),
        ({"queue": {"max_attempts": 1}}, "AI settings are invalid"),
    ],
)
def test_tracked_settings_reject_unsupported_v1_contracts(
    tmp_path: Path, change: dict[str, object], message: str
) -> None:
    config = tmp_path / "ai.yaml"
    write_yaml(config, ai_values())

    with pytest.raises(SettingsError, match=message):
        load_ai_settings(config, change)


@pytest.mark.parametrize(
    "secret_key",
    ["api_key", "OPENAI-API-KEY", "access_token", "client_secret"],
)
def test_tracked_yaml_rejects_secret_like_keys(
    tmp_path: Path, secret_key: str
) -> None:
    config = tmp_path / "ai.yaml"
    values = ai_values()
    values["provider"][secret_key] = "do-not-track"
    write_yaml(config, values)

    with pytest.raises(SettingsError, match="Secret-like key"):
        load_ai_settings(config)


def test_tracked_settings_reject_missing_and_corrupt_yaml(
    tmp_path: Path,
) -> None:
    with pytest.raises(SettingsError, match="Cannot read YAML"):
        load_ai_settings(tmp_path / "missing.yaml")

    config = tmp_path / "ai.yaml"
    config.write_text("provider: [unterminated", encoding="utf-8")
    with pytest.raises(SettingsError, match="YAML file is invalid"):
        load_ai_settings(config)


def test_private_inputs_validate_paths_types_layout_and_hashes(
    tmp_path: Path,
) -> None:
    profile_root = write_private_inputs(tmp_path)

    inputs = validate_private_inputs(tmp_path)

    assert inputs.profile_path == profile_root / "profile.md"
    assert inputs.layout.page.size == "A4"
    assert inputs.layout.output.filename == "Tailored CV.docx"
    assert len(inputs.profile_sha256) == 64
    assert len(inputs.template_sha256) == 64
    assert len(inputs.layout_sha256) == 64


def test_private_input_validators_compose_without_changing_hashes(
    tmp_path: Path,
) -> None:
    profile_root = write_private_inputs(tmp_path)

    profile = validate_profile_input(tmp_path)
    template = validate_template_input(tmp_path)
    layout_path, layout = validate_layout_input(tmp_path)
    aggregate = validate_private_inputs(tmp_path)

    assert profile == profile_root / "profile.md"
    assert template == profile_root / "cv-template.docx"
    assert layout_path == profile_root / "cv-layout.yaml"
    assert layout == aggregate.layout
    assert aggregate.profile_sha256
    assert aggregate.template_sha256
    assert aggregate.layout_sha256


@pytest.mark.parametrize(
    "style_name", ["Normal", "Title", "Heading 1", "List Bullet"]
)
def test_template_input_rejects_missing_required_styles(
    tmp_path: Path, style_name: str
) -> None:
    profile_root = write_private_inputs(tmp_path)
    template_path = profile_root / "cv-template.docx"
    document = Document(str(template_path))
    style = document.styles[style_name]
    style.element.getparent().remove(style.element)
    document.save(str(template_path))

    with pytest.raises(
        SettingsError, match=f"missing required style: {style_name}"
    ):
        validate_template_input(tmp_path)


def test_template_input_rejects_required_non_paragraph_style(
    tmp_path: Path,
) -> None:
    profile_root = write_private_inputs(tmp_path)
    template_path = profile_root / "cv-template.docx"
    document = Document(str(template_path))
    document.styles["Title"].element.set(qn("w:type"), "character")
    document.save(str(template_path))

    with pytest.raises(SettingsError, match="must be a paragraph style: Title"):
        validate_template_input(tmp_path)


def test_private_inputs_reject_missing_corrupt_and_oversized_files(
    tmp_path: Path,
) -> None:
    profile_root = write_private_inputs(tmp_path)
    (profile_root / "cv-template.docx").unlink()
    with pytest.raises(SettingsError, match="Missing or invalid private"):
        validate_private_inputs(tmp_path)

    (profile_root / "cv-template.docx").write_bytes(b"not a docx")
    with pytest.raises(SettingsError, match="not a valid DOCX"):
        validate_private_inputs(tmp_path)

    document = Document()
    document.add_paragraph("Template")
    document.save(str(profile_root / "cv-template.docx"))
    (profile_root / "profile.md").write_bytes(b"x" * (MAX_PROFILE_BYTES + 1))
    with pytest.raises(SettingsError, match="must be between"):
        validate_private_inputs(tmp_path)


def test_private_inputs_reject_invalid_layout_and_outside_override(
    tmp_path: Path,
) -> None:
    profile_root = write_private_inputs(tmp_path)
    invalid_layout = layout_values()
    invalid_layout["page"]["margins"]["top"] = 2
    write_yaml(profile_root / "cv-layout.yaml", invalid_layout)

    with pytest.raises(SettingsError, match="CV layout is invalid"):
        validate_private_inputs(tmp_path)

    outside = tmp_path / "outside.md"
    outside.write_text("# Outside", encoding="utf-8")
    with pytest.raises(SettingsError, match="outside"):
        validate_private_inputs(tmp_path, profile_path=outside)


def test_private_inputs_reject_invalid_layout_relationship(
    tmp_path: Path,
) -> None:
    profile_root = write_private_inputs(tmp_path)
    invalid_layout = layout_values()
    invalid_layout["fonts"]["body_size_pt"] = 12
    invalid_layout["fonts"]["heading_size_pt"] = 11
    write_yaml(profile_root / "cv-layout.yaml", invalid_layout)

    with pytest.raises(SettingsError, match="heading font"):
        validate_private_inputs(tmp_path)


def test_private_inputs_reject_symlinks(tmp_path: Path) -> None:
    profile_root = write_private_inputs(tmp_path)
    real_profile = profile_root / "real.md"
    real_profile.write_text("# Real", encoding="utf-8")
    (profile_root / "profile.md").unlink()
    (profile_root / "profile.md").symlink_to(real_profile)

    with pytest.raises(SettingsError, match="must not use symlinks"):
        validate_private_inputs(tmp_path)
