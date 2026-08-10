from __future__ import annotations

import hashlib
import json
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Any, Literal
from xml.etree import ElementTree

import pytest
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.shared import Inches, Mm

from app.ai.schema_models import CvContent
from app.artefacts import ArtefactStore, UnsafeArtefactPathError
from app.documents.word_writer import (
    MAX_CANDIDATE_BYTES,
    WordWriter,
    WordWriterError,
    _sanitize_package,
    candidate_filename,
)
from app.settings import CvLayoutSettings


def test_documents_public_boundary_only_exposes_docx_writer_contract() -> None:
    from app import documents

    assert documents.__all__ == [
        "CandidateDocument",
        "WordWriter",
        "WordWriterError",
    ]
    assert {
        "CandidateDocument",
        "WordWriter",
        "WordWriterError",
    } == set(documents.__all__)
    for removed_name in (
        "DocumentVerifier",
        "FontInventory",
        "ProcessRunner",
        "WordVerificationRequest",
        "WordVerificationResult",
        "WordVerifier",
    ):
        assert not hasattr(documents, removed_name)


def layout() -> CvLayoutSettings:
    return CvLayoutSettings.model_validate(
        {
            "page": {
                "size": "A4",
                "orientation": "portrait",
                "margins": {
                    "top": 0.5,
                    "right": 0.55,
                    "bottom": 0.5,
                    "left": 0.55,
                },
            },
            "fonts": {
                "body": "Aptos",
                "headings": "Aptos Display",
                "body_size_pt": 9.5,
                "heading_size_pt": 12,
            },
            "spacing": {
                "line": 1.0,
                "paragraph_after_pt": 3,
                "section_after_pt": 6,
            },
            "styles": {
                "heading_color": "1F2937",
                "accent_color": "2563EB",
                "bullet_indent_inches": 0.2,
            },
            "output": {"filename": "Tailored CV.docx"},
        }
    )


def content(**changes: Any) -> CvContent:
    values: dict[str, Any] = {
        "schema_version": "v1",
        "identity": {
            "full_name": {
                "text": "Avery Morgan",
                "profile_block_ids": ["profile-0001"],
            },
            "professional_title": {
                "text": "Software Engineer",
                "profile_block_ids": ["profile-0002"],
            },
            "email": {
                "text": "avery@example.test",
                "profile_block_ids": ["profile-0002"],
            },
            "phone": None,
            "location": None,
            "website": {
                "text": "https://example.test/avery",
                "profile_block_ids": ["profile-0002"],
            },
        },
        "sections": [
            {
                "section_id": "section-01",
                "kind": "summary",
                "heading": "Profile",
                "claims": [
                    {
                        "claim_id": "claim-001",
                        "text": "Builds typed Python services.",
                        "profile_block_ids": ["profile-0002"],
                        "requirement_ids": ["req-001"],
                    }
                ],
            },
            {
                "section_id": "section-02",
                "kind": "skills",
                "heading": "Skills",
                "claims": [
                    {
                        "claim_id": "claim-002",
                        "text": "Python, FastAPI",
                        "profile_block_ids": ["profile-0002"],
                        "requirement_ids": ["req-001"],
                    }
                ],
            },
        ],
    }
    values.update(changes)
    return CvContent.model_validate_json(json.dumps(values))


def template(path: Path) -> None:
    document = Document()
    document.add_paragraph("Private placeholder that must be removed.")
    document.sections[0].header.add_paragraph("Private header placeholder")
    document.core_properties.author = "Template Author"
    document.save(str(path))


def background_template(
    path: Path, relationship_kind: Literal["external", "internal", "none"]
) -> None:
    """Create a template with a document-level private background."""
    template(path)
    source = BytesIO(path.read_bytes())
    destination = BytesIO()
    with zipfile.ZipFile(source) as archive:
        members = {
            member.filename: archive.read(member.filename)
            for member in archive.infolist()
        }

    document_root = ElementTree.fromstring(members["word/document.xml"])
    background = ElementTree.Element(qn("w:background"))
    if relationship_kind != "none":
        background.set(qn("r:id"), "rIdPrivateBackground")
    document_root.insert(0, background)
    members["word/document.xml"] = ElementTree.tostring(
        document_root,
        encoding="utf-8",
        xml_declaration=True,
    )

    relationships = ElementTree.fromstring(
        members["word/_rels/document.xml.rels"]
    )
    if relationship_kind == "internal":
        ElementTree.SubElement(
            relationships,
            f"{{{PACKAGE_RELATIONSHIP_NAMESPACE}}}Relationship",
            {
                "Id": "rIdPrivateBackground",
                "Type": f"{OFFICE_RELATIONSHIP_NAMESPACE}/image",
                "Target": "media/private-background.png",
            },
        )
        members["word/media/private-background.png"] = (
            b"PRIVATE BACKGROUND SENTINEL"
        )
        content_types = ElementTree.fromstring(members["[Content_Types].xml"])
        ElementTree.SubElement(
            content_types,
            "{http://schemas.openxmlformats.org/package/2006/content-types}"
            "Override",
            {
                "PartName": "/word/media/private-background.png",
                "ContentType": "image/png",
            },
        )
        members["[Content_Types].xml"] = ElementTree.tostring(
            content_types,
            encoding="utf-8",
            xml_declaration=True,
        )
    elif relationship_kind == "external":
        ElementTree.SubElement(
            relationships,
            f"{{{PACKAGE_RELATIONSHIP_NAMESPACE}}}Relationship",
            {
                "Id": "rIdPrivateBackground",
                "Type": f"{OFFICE_RELATIONSHIP_NAMESPACE}/image",
                "Target": "https://private-background.example.test/image.png",
                "TargetMode": "External",
            },
        )
    members["word/_rels/document.xml.rels"] = ElementTree.tostring(
        relationships,
        encoding="utf-8",
        xml_declaration=True,
    )
    with zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    path.write_bytes(destination.getvalue())


def private_template_content(path: Path) -> None:
    document = Document()
    document.add_paragraph("Private body placeholder")
    header_table = document.sections[0].header.add_table(1, 1, Inches(6))
    header_table.cell(0, 0).text = "Private header table"
    footer_table = document.sections[0].footer.add_table(1, 1, Inches(6))
    footer_table.cell(0, 0).text = "Private footer table"
    document.save(str(path))

    relationship_namespace = (
        "http://schemas.openxmlformats.org/package/2006/relationships"
    )
    relationships = ElementTree.Element(
        f"{{{relationship_namespace}}}Relationships"
    )
    ElementTree.SubElement(
        relationships,
        f"{{{relationship_namespace}}}Relationship",
        {
            "Id": "rIdPrivateImage",
            "Type": (
                "http://schemas.openxmlformats.org/officeDocument/2006/"
                "relationships/image"
            ),
            "Target": "media/private-image.png",
        },
    )
    ElementTree.SubElement(
        relationships,
        f"{{{relationship_namespace}}}Relationship",
        {
            "Id": "rIdPrivateLink",
            "Type": (
                "http://schemas.openxmlformats.org/officeDocument/2006/"
                "relationships/hyperlink"
            ),
            "Target": "https://private.example.test",
            "TargetMode": "External",
        },
    )
    ElementTree.SubElement(
        relationships,
        f"{{{relationship_namespace}}}Relationship",
        {
            "Id": "rIdPrivateObject",
            "Type": (
                "http://schemas.openxmlformats.org/officeDocument/2006/"
                "relationships/oleObject"
            ),
            "Target": "embeddings/private-object.bin",
        },
    )
    ElementTree.SubElement(
        relationships,
        f"{{{relationship_namespace}}}Relationship",
        {
            "Id": "rIdPrivateControl",
            "Type": (
                "http://schemas.openxmlformats.org/officeDocument/2006/"
                "relationships/control"
            ),
            "Target": "activeX/private-control.xml",
        },
    )
    source = BytesIO(path.read_bytes())
    destination = BytesIO()
    with zipfile.ZipFile(source) as archive:
        members = {
            member.filename: archive.read(member.filename)
            for member in archive.infolist()
        }
    members["word/_rels/header1.xml.rels"] = ElementTree.tostring(
        relationships,
        encoding="utf-8",
        xml_declaration=True,
    )
    members["word/media/private-image.png"] = b"private image bytes"
    members["word/embeddings/private-object.bin"] = b"private object bytes"
    members["word/activeX/private-control.xml"] = b"<private-control/>"
    nested_relationships = ElementTree.Element(
        f"{{{relationship_namespace}}}Relationships"
    )
    for relationship_id, target in (
        ("rIdPrivateBinary", "../embeddings/private-active.bin"),
        ("rIdPrivatePreview", "../media/private-preview.png"),
    ):
        ElementTree.SubElement(
            nested_relationships,
            f"{{{relationship_namespace}}}Relationship",
            {
                "Id": relationship_id,
                "Type": (
                    "http://schemas.openxmlformats.org/officeDocument/"
                    "2006/relationships/image"
                ),
                "Target": target,
            },
        )
    members["word/activeX/_rels/private-control.xml.rels"] = (
        ElementTree.tostring(
            nested_relationships,
            encoding="utf-8",
            xml_declaration=True,
        )
    )
    members["word/embeddings/private-active.bin"] = b"private active bytes"
    members["word/media/private-preview.png"] = b"private preview bytes"
    root_relationships = ElementTree.fromstring(members["_rels/.rels"])
    for relationship_id, relationship_type, target in (
        (
            "rIdPrivateCustomProperties",
            "custom-properties",
            "docProps/custom.xml",
        ),
        (
            "rIdPrivateThumbnail",
            "thumbnail",
            "docProps/thumbnail.jpeg",
        ),
    ):
        relationship_namespace_prefix = (
            "http://schemas.openxmlformats.org/officeDocument/2006/"
            "relationships"
        )
        if relationship_type == "thumbnail":
            relationship_namespace_prefix = (
                "http://schemas.openxmlformats.org/package/2006/"
                "relationships/metadata"
            )
        ElementTree.SubElement(
            root_relationships,
            f"{{{relationship_namespace}}}Relationship",
            {
                "Id": relationship_id,
                "Type": (
                    f"{relationship_namespace_prefix}/{relationship_type}"
                ),
                "Target": target,
            },
        )
    members["_rels/.rels"] = ElementTree.tostring(
        root_relationships,
        encoding="utf-8",
        xml_declaration=True,
    )
    document_relationships = ElementTree.fromstring(
        members["word/_rels/document.xml.rels"]
    )
    for relationship_id, relationship_type, target in (
        ("rIdPrivateComments", "comments", "comments.xml"),
        ("rIdPrivateFootnotes", "footnotes", "footnotes.xml"),
        (
            "rIdPrivateGlossary",
            "glossaryDocument",
            "glossary/document.xml",
        ),
    ):
        ElementTree.SubElement(
            document_relationships,
            f"{{{relationship_namespace}}}Relationship",
            {
                "Id": relationship_id,
                "Type": (
                    "http://schemas.openxmlformats.org/officeDocument/2006/"
                    f"relationships/{relationship_type}"
                ),
                "Target": target,
            },
        )
    members["word/_rels/document.xml.rels"] = ElementTree.tostring(
        document_relationships,
        encoding="utf-8",
        xml_declaration=True,
    )
    members["docProps/custom.xml"] = (
        b"<custom>CONFIDENTIAL CUSTOM PROPERTY</custom>"
    )
    members["docProps/thumbnail.jpeg"] = b"CONFIDENTIAL THUMBNAIL"
    members["word/comments.xml"] = b"<comments>CONFIDENTIAL COMMENT</comments>"
    members["word/footnotes.xml"] = (
        b"<footnotes>CONFIDENTIAL FOOTNOTE</footnotes>"
    )
    members["word/glossary/document.xml"] = (
        b"<glossary>CONFIDENTIAL GLOSSARY</glossary>"
    )
    members["word/_rels/comments.xml.rels"] = _relationship_xml(
        (
            (
                "rIdPrivateCommentImage",
                (
                    "http://schemas.openxmlformats.org/officeDocument/2006/"
                    "relationships/image"
                ),
                "media/private-comment.png",
                None,
            ),
        )
    )
    members["word/media/private-comment.png"] = (
        b"CONFIDENTIAL COMMENT DESCENDANT"
    )
    content_types = ElementTree.fromstring(members["[Content_Types].xml"])
    ElementTree.SubElement(
        content_types,
        "{http://schemas.openxmlformats.org/package/2006/content-types}Default",
        {"Extension": "png", "ContentType": "image/png"},
    )
    ElementTree.SubElement(
        content_types,
        "{http://schemas.openxmlformats.org/package/2006/content-types}Override",
        {
            "PartName": "/word/embeddings/private-object.bin",
            "ContentType": "application/vnd.openxmlformats-officedocument."
            "oleObject",
        },
    )
    for part_name, content_type in (
        (
            "/word/activeX/private-control.xml",
            "application/vnd.ms-office.activeX+xml",
        ),
        (
            "/word/embeddings/private-active.bin",
            "application/vnd.ms-office.activeX",
        ),
        (
            "/docProps/custom.xml",
            (
                "application/vnd.openxmlformats-officedocument."
                "custom-properties+xml"
            ),
        ),
        ("/docProps/thumbnail.jpeg", "image/jpeg"),
        (
            "/word/comments.xml",
            (
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.comments+xml"
            ),
        ),
        (
            "/word/footnotes.xml",
            (
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.footnotes+xml"
            ),
        ),
        (
            "/word/glossary/document.xml",
            (
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document.glossary+xml"
            ),
        ),
        ("/word/media/private-comment.png", "image/png"),
    ):
        ElementTree.SubElement(
            content_types,
            "{http://schemas.openxmlformats.org/package/2006/content-types}"
            "Override",
            {"PartName": part_name, "ContentType": content_type},
        )
    members["[Content_Types].xml"] = ElementTree.tostring(
        content_types,
        encoding="utf-8",
        xml_declaration=True,
    )
    settings = members["word/settings.xml"].decode("utf-8")
    members["word/settings.xml"] = settings.replace(
        "</w:settings>",
        (
            '<w:zoom w:percent="90"/>'
            '<w:attachedTemplate r:id="rIdPrivateTemplate"/>'
            '<w:docVars><w:docVar w:name="Client" '
            'w:val="CONFIDENTIAL DOCUMENT VARIABLE"/></w:docVars>'
            '<w:mailMerge><w:dataSource r:id="rIdPrivateDataSource"/>'
            "</w:mailMerge>"
            '<mc:AlternateContent><mc:Choice Requires="w14">'
            '<w:zoom w:percent="85"/>'
            '<w:attachedTemplate r:id="rIdChoiceTemplate"/>'
            '<w:docVars><w:docVar w:name="ChoiceClient" '
            'w:val="CONFIDENTIAL CHOICE VARIABLE"/></w:docVars>'
            '<w:mailMerge><w:dataSource r:id="rIdChoiceDataSource"/>'
            "</w:mailMerge></mc:Choice><mc:Fallback>"
            "<w:compat><w:attachedTemplate "
            'r:id="rIdFallbackTemplate"/><w:docVars>'
            '<w:docVar w:name="FallbackClient" '
            'w:val="CONFIDENTIAL FALLBACK VARIABLE"/></w:docVars>'
            "<w:mailMerge><w:dataSource "
            'r:id="rIdFallbackDataSource"/></w:mailMerge>'
            '<w:compatSetting w:name="safe" w:uri="urn:safe" '
            'w:val="1"/></w:compat></mc:Fallback>'
            "</mc:AlternateContent></w:settings>"
        ),
    ).encode("utf-8")
    settings_relationships = ElementTree.Element(
        f"{{{relationship_namespace}}}Relationships"
    )
    ElementTree.SubElement(
        settings_relationships,
        f"{{{relationship_namespace}}}Relationship",
        {
            "Id": "rIdPrivateTemplate",
            "Type": (
                "http://schemas.openxmlformats.org/officeDocument/2006/"
                "relationships/attachedTemplate"
            ),
            "Target": "https://private-template.example.test",
            "TargetMode": "External",
        },
    )
    for relationship_id, relationship_kind, target in (
        (
            "rIdChoiceTemplate",
            "attachedTemplate",
            "https://choice-template.example.test",
        ),
        (
            "rIdChoiceDataSource",
            "mailMergeSource",
            "file:///private/choice-client-list.xlsx",
        ),
        (
            "rIdFallbackTemplate",
            "attachedTemplate",
            "https://fallback-template.example.test",
        ),
        (
            "rIdFallbackDataSource",
            "mailMergeSource",
            "file:///private/fallback-client-list.xlsx",
        ),
    ):
        ElementTree.SubElement(
            settings_relationships,
            f"{{{relationship_namespace}}}Relationship",
            {
                "Id": relationship_id,
                "Type": (
                    "http://schemas.openxmlformats.org/officeDocument/2006/"
                    f"relationships/{relationship_kind}"
                ),
                "Target": target,
                "TargetMode": "External",
            },
        )
    ElementTree.SubElement(
        settings_relationships,
        f"{{{relationship_namespace}}}Relationship",
        {
            "Id": "rIdPrivateDataSource",
            "Type": (
                "http://schemas.openxmlformats.org/officeDocument/2006/"
                "relationships/mailMergeSource"
            ),
            "Target": "file:///private/client-list.xlsx",
            "TargetMode": "External",
        },
    )
    members["word/_rels/settings.xml.rels"] = ElementTree.tostring(
        settings_relationships,
        encoding="utf-8",
        xml_declaration=True,
    )
    core = ElementTree.fromstring(members["docProps/core.xml"])
    ElementTree.SubElement(
        core,
        (
            "{http://schemas.openxmlformats.org/package/2006/"
            "metadata/core-properties}contentStatus"
        ),
    ).text = "CONFIDENTIAL REVIEW STATUS"
    ElementTree.SubElement(
        core, "{urn:jobhunter:test-private}metadata"
    ).text = "CONFIDENTIAL UNKNOWN CORE PROPERTY"
    members["docProps/core.xml"] = ElementTree.tostring(
        core,
        encoding="utf-8",
        xml_declaration=True,
    )
    with zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    path.write_bytes(destination.getvalue())


def private_section_metadata_template(path: Path) -> None:
    """Add private section revision and printer settings to a template."""
    private_template_content(path)
    source = BytesIO(path.read_bytes())
    destination = BytesIO()
    with zipfile.ZipFile(source) as archive:
        members = {
            member.filename: archive.read(member.filename)
            for member in archive.infolist()
        }

    document_root = ElementTree.fromstring(members["word/document.xml"])
    section_properties = document_root.find(
        f".//{qn('w:body')}/{qn('w:sectPr')}"
    )
    assert section_properties is not None
    section_properties.set(qn("w:rsidR"), "PRIVATE-RSID")
    section_change = ElementTree.SubElement(
        section_properties,
        qn("w:sectPrChange"),
        {
            qn("w:id"): "7",
            qn("w:author"): "Private Reviewer",
            qn("w:date"): "2026-08-09T12:00:00Z",
        },
    )
    ElementTree.SubElement(
        section_change,
        qn("w:pgSz"),
        {qn("w:w"): "1", qn("w:h"): "1"},
    )
    printer_settings = ElementTree.SubElement(
        section_properties, qn("w:printerSettings")
    )
    printer_settings.set(qn("r:id"), "rIdPrivatePrinterSettings")
    members["word/document.xml"] = ElementTree.tostring(
        document_root,
        encoding="utf-8",
        xml_declaration=True,
    )

    document_relationships = ElementTree.fromstring(
        members["word/_rels/document.xml.rels"]
    )
    ElementTree.SubElement(
        document_relationships,
        f"{{{PACKAGE_RELATIONSHIP_NAMESPACE}}}Relationship",
        {
            "Id": "rIdPrivatePrinterSettings",
            "Type": (f"{OFFICE_RELATIONSHIP_NAMESPACE}/printerSettings"),
            "Target": "printerSettings1.bin",
        },
    )
    members["word/_rels/document.xml.rels"] = ElementTree.tostring(
        document_relationships,
        encoding="utf-8",
        xml_declaration=True,
    )
    members["word/printerSettings1.bin"] = b"PRIVATE PRINTER SETTINGS"
    content_types = ElementTree.fromstring(members["[Content_Types].xml"])
    ElementTree.SubElement(
        content_types,
        "{http://schemas.openxmlformats.org/package/2006/content-types}"
        "Override",
        {
            "PartName": "/word/printerSettings1.bin",
            "ContentType": (
                "application/vnd.openxmlformats-officedocument.wordprocessingml."
                "printerSettings"
            ),
        },
    )
    members["[Content_Types].xml"] = ElementTree.tostring(
        content_types,
        encoding="utf-8",
        xml_declaration=True,
    )
    with zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    path.write_bytes(destination.getvalue())


PACKAGE_RELATIONSHIP_NAMESPACE = (
    "http://schemas.openxmlformats.org/package/2006/relationships"
)
OFFICE_RELATIONSHIP_NAMESPACE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
)
PACKAGE_METADATA_RELATIONSHIP_NAMESPACE = (
    "http://schemas.openxmlformats.org/package/2006/relationships/metadata"
)
VML_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.vmlDrawing"
WORD_DOCUMENT_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument."
    "wordprocessingml.document.main+xml"
)


def _relationship_xml(
    relationships: tuple[tuple[str, str, str, str | None], ...],
) -> bytes:
    root = ElementTree.Element(
        f"{{{PACKAGE_RELATIONSHIP_NAMESPACE}}}Relationships"
    )
    for (
        relationship_id,
        relationship_type,
        target,
        target_mode,
    ) in relationships:
        attributes = {
            "Id": relationship_id,
            "Type": relationship_type,
            "Target": target,
        }
        if target_mode is not None:
            attributes["TargetMode"] = target_mode
        ElementTree.SubElement(
            root,
            f"{{{PACKAGE_RELATIONSHIP_NAMESPACE}}}Relationship",
            attributes,
        )
    return ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)


def _sanitize_members(members: dict[str, bytes]) -> dict[str, bytes]:
    source = BytesIO()
    with zipfile.ZipFile(source, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    with zipfile.ZipFile(
        BytesIO(_sanitize_package(source.getvalue()))
    ) as archive:
        return {
            member.filename: archive.read(member.filename)
            for member in archive.infolist()
        }


def _content_types(*part_names: str) -> bytes:
    root = ElementTree.Element(
        "{http://schemas.openxmlformats.org/package/2006/content-types}Types"
    )
    for part_name in part_names:
        ElementTree.SubElement(
            root,
            "{http://schemas.openxmlformats.org/package/2006/content-types}"
            "Override",
            {"PartName": f"/{part_name}", "ContentType": "application/test"},
        )
    return ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)


def _typed_content_types(
    defaults: tuple[tuple[str, str], ...],
    overrides: tuple[tuple[str, str], ...],
) -> bytes:
    namespace = "http://schemas.openxmlformats.org/package/2006/content-types"
    root = ElementTree.Element(f"{{{namespace}}}Types")
    for extension, content_type in defaults:
        ElementTree.SubElement(
            root,
            f"{{{namespace}}}Default",
            {"Extension": extension, "ContentType": content_type},
        )
    for part_name, content_type in overrides:
        ElementTree.SubElement(
            root,
            f"{{{namespace}}}Override",
            {"PartName": f"/{part_name}", "ContentType": content_type},
        )
    return ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)


def _relationship_targets(data: bytes) -> set[str]:
    return {
        relationship.get("Target", "")
        for relationship in ElementTree.fromstring(data)
    }


def _relationship_ids(data: bytes) -> set[str]:
    return {
        relationship.get("Id", "")
        for relationship in ElementTree.fromstring(data)
    }


def test_sanitizer_prunes_private_parts_and_keeps_rendering_parts() -> None:
    document_xml = f"""<w:document
 xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
 xmlns:r="{OFFICE_RELATIONSHIP_NAMESPACE}">
 <w:hyperlink r:id="rIdHyperlink"/>
 </w:document>""".encode()
    rendering_parts = {
        "word/styles.xml": b"<styles>safe styles</styles>",
        "word/settings.xml": b"<settings>safe settings</settings>",
        "word/numbering.xml": b"<numbering>safe numbering</numbering>",
        "word/fontTable.xml": b"<fonts>safe fonts</fonts>",
        "word/webSettings.xml": b"<web>safe web settings</web>",
        "word/theme/theme1.xml": b"<theme>safe theme</theme>",
    }
    private_parts = {
        "docProps/custom.xml": b"<custom>CONFIDENTIAL CUSTOM</custom>",
        "docProps/app.xml": b"<app>CONFIDENTIAL EXTENDED</app>",
        "docProps/thumbnail.jpeg": b"CONFIDENTIAL THUMBNAIL",
        "word/comments.xml": b"<comments>CONFIDENTIAL COMMENT</comments>",
        "word/footnotes.xml": b"<footnotes>CONFIDENTIAL FOOTNOTE</footnotes>",
        "word/glossary/document.xml": (
            b"<glossary>CONFIDENTIAL GLOSSARY</glossary>"
        ),
        "word/media/comment-preview.png": b"CONFIDENTIAL DESCENDANT",
    }
    part_names = (
        "docProps/core.xml",
        "word/document.xml",
        *rendering_parts,
        *private_parts,
    )
    members = {
        "[Content_Types].xml": _content_types(*part_names),
        "_rels/.rels": _relationship_xml(
            (
                (
                    "rIdDocument",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/officeDocument",
                    "word/document.xml",
                    None,
                ),
                (
                    "rIdCore",
                    (
                        f"{PACKAGE_METADATA_RELATIONSHIP_NAMESPACE}/"
                        "core-properties"
                    ),
                    "docProps/core.xml",
                    None,
                ),
                (
                    "rIdCustom",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/custom-properties",
                    "docProps/custom.xml",
                    None,
                ),
                (
                    "rIdExtended",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/extended-properties",
                    "docProps/app.xml",
                    None,
                ),
                (
                    "rIdThumbnail",
                    f"{PACKAGE_METADATA_RELATIONSHIP_NAMESPACE}/thumbnail",
                    "docProps/thumbnail.jpeg",
                    None,
                ),
            )
        ),
        "docProps/core.xml": b"<core>safe core properties</core>",
        "word/document.xml": document_xml,
        "word/_rels/document.xml.rels": _relationship_xml(
            (
                *(
                    (
                        f"rId{index}",
                        f"{OFFICE_RELATIONSHIP_NAMESPACE}/{kind}",
                        target,
                        None,
                    )
                    for index, (kind, target) in enumerate(
                        (
                            ("styles", "styles.xml"),
                            ("settings", "settings.xml"),
                            ("numbering", "numbering.xml"),
                            ("fontTable", "fontTable.xml"),
                            ("webSettings", "webSettings.xml"),
                            ("theme", "theme/theme1.xml"),
                            ("comments", "comments.xml"),
                            ("footnotes", "footnotes.xml"),
                            (
                                "glossaryDocument",
                                "glossary/document.xml",
                            ),
                        ),
                        start=1,
                    )
                ),
                (
                    "rIdHyperlink",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/hyperlink",
                    "https://candidate.example.test",
                    "External",
                ),
            )
        ),
        "word/_rels/comments.xml.rels": _relationship_xml(
            (
                (
                    "rIdCommentPreview",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/image",
                    "media/comment-preview.png",
                    None,
                ),
            )
        ),
        **rendering_parts,
        **private_parts,
    }

    sanitized = _sanitize_members(members)

    assert "docProps/core.xml" in sanitized
    assert set(rendering_parts) <= sanitized.keys()
    assert _relationship_targets(sanitized["word/_rels/document.xml.rels"]) == {
        "styles.xml",
        "settings.xml",
        "numbering.xml",
        "fontTable.xml",
        "webSettings.xml",
        "theme/theme1.xml",
        "https://candidate.example.test",
    }
    for part_name in private_parts:
        assert part_name not in sanitized
    assert "word/_rels/comments.xml.rels" not in sanitized
    assert b"CONFIDENTIAL" not in b"".join(sanitized.values())
    content_types = sanitized["[Content_Types].xml"]
    for part_name in private_parts:
        assert part_name.encode() not in content_types


def test_sanitizer_keeps_microsoft_effects_aware_styles() -> None:
    effects_relationship_type = (
        "http://schemas.microsoft.com/office/2007/relationships/"
        "stylesWithEffects"
    )
    effects_content_type = "application/vnd.ms-word.stylesWithEffects+xml"
    members = {
        "[Content_Types].xml": _typed_content_types(
            defaults=(),
            overrides=(
                ("word/document.xml", WORD_DOCUMENT_CONTENT_TYPE),
                ("word/stylesWithEffects.xml", effects_content_type),
                ("word/orphanStylesWithEffects.xml", effects_content_type),
                ("word/comments.xml", "application/test-comments+xml"),
            ),
        ),
        "_rels/.rels": _relationship_xml(
            (
                (
                    "rIdDocument",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/officeDocument",
                    "word/document.xml",
                    None,
                ),
            )
        ),
        "word/document.xml": b"<document/>",
        "word/_rels/document.xml.rels": _relationship_xml(
            (
                (
                    "rIdEffectsStyles",
                    effects_relationship_type,
                    "stylesWithEffects.xml",
                    None,
                ),
                (
                    "rIdPrivateComments",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/comments",
                    "comments.xml",
                    None,
                ),
            )
        ),
        "word/stylesWithEffects.xml": b"<styles>safe effects</styles>",
        "word/orphanStylesWithEffects.xml": b"<styles>orphan</styles>",
        "word/comments.xml": b"<comments>CONFIDENTIAL</comments>",
    }

    sanitized = _sanitize_members(members)

    relationship_name = "word/_rels/document.xml.rels"
    assert relationship_name in sanitized
    relationships = tuple(ElementTree.fromstring(sanitized[relationship_name]))
    assert len(relationships) == 1
    assert relationships[0].attrib == {
        "Id": "rIdEffectsStyles",
        "Type": effects_relationship_type,
        "Target": "stylesWithEffects.xml",
    }
    assert sanitized["word/stylesWithEffects.xml"] == (
        b"<styles>safe effects</styles>"
    )
    assert "word/orphanStylesWithEffects.xml" not in sanitized
    assert "word/comments.xml" not in sanitized
    content_types = sanitized["[Content_Types].xml"]
    assert b"/word/stylesWithEffects.xml" in content_types
    assert effects_content_type.encode() in content_types
    assert b"orphanStylesWithEffects.xml" not in content_types
    assert b"comments.xml" not in content_types
    assert b"CONFIDENTIAL" not in b"".join(sanitized.values())


def test_sanitizer_keeps_live_relationships_from_vml_owner() -> None:
    document_xml = f"""<w:document
 xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
 xmlns:r="{OFFICE_RELATIONSHIP_NAMESPACE}">
 <w:object r:id="rIdVmlDrawing"/>
 </w:document>""".encode()
    vml_xml = f"""<v:shape
 xmlns:v="urn:schemas-microsoft-com:vml"
 xmlns:r="{OFFICE_RELATIONSHIP_NAMESPACE}">
 <v:imagedata r:id="rIdImage"/>
 <v:fill r:embed="rIdEmbeddedImage" r:link="rIdLinkedImage"/>
 </v:shape>""".encode()
    members = {
        "[Content_Types].xml": _typed_content_types(
            defaults=(("vml", VML_CONTENT_TYPE),),
            overrides=(
                ("word/document.xml", WORD_DOCUMENT_CONTENT_TYPE),
                ("word/media/live-id.png", "image/png"),
                ("word/media/live-embed.png", "image/png"),
                ("word/media/private.png", "image/png"),
            ),
        ),
        "_rels/.rels": _relationship_xml(
            (
                (
                    "rIdDocument",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/officeDocument",
                    "word/document.xml",
                    None,
                ),
            )
        ),
        "word/document.xml": document_xml,
        "word/_rels/document.xml.rels": _relationship_xml(
            (
                (
                    "rIdVmlDrawing",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/vmlDrawing",
                    "drawings/live.vml",
                    None,
                ),
            )
        ),
        "word/drawings/live.vml": vml_xml,
        "word/drawings/_rels/live.vml.rels": _relationship_xml(
            (
                (
                    "rIdImage",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/image",
                    "../media/live-id.png",
                    None,
                ),
                (
                    "rIdEmbeddedImage",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/image",
                    "../media/live-embed.png",
                    None,
                ),
                (
                    "rIdLinkedImage",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/image",
                    "https://images.example.test/linked.png",
                    "External",
                ),
                (
                    "rIdUnusedImage",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/image",
                    "../media/private.png",
                    None,
                ),
            )
        ),
        "word/media/live-id.png": b"live id image",
        "word/media/live-embed.png": b"live embedded image",
        "word/media/private.png": b"CONFIDENTIAL UNUSED VML IMAGE",
    }

    sanitized = _sanitize_members(members)

    vml_relationships = sanitized["word/drawings/_rels/live.vml.rels"]
    assert _relationship_ids(vml_relationships) == {
        "rIdImage",
        "rIdEmbeddedImage",
        "rIdLinkedImage",
    }
    assert _relationship_targets(vml_relationships) == {
        "../media/live-id.png",
        "../media/live-embed.png",
        "https://images.example.test/linked.png",
    }
    assert "word/media/live-id.png" in sanitized
    assert "word/media/live-embed.png" in sanitized
    assert "word/media/private.png" not in sanitized
    assert b"private.png" not in sanitized["[Content_Types].xml"]
    assert b"CONFIDENTIAL" not in b"".join(sanitized.values())


@pytest.mark.parametrize(
    ("target", "part_name"),
    (
        ("/word/media/live.png", "word/media/live.png"),
        ("media/live.png", "word/media/live.png"),
        ("./media/live.png", "word/media/live.png"),
        ("../live.png", "live.png"),
    ),
)
def test_sanitizer_resolves_live_internal_relationship_targets(
    target: str,
    part_name: str,
) -> None:
    document_xml = f"""<w:document
 xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
 xmlns:r="{OFFICE_RELATIONSHIP_NAMESPACE}">
 <w:drawing r:id="rIdLiveImage"/>
 </w:document>""".encode()
    members = {
        "[Content_Types].xml": _typed_content_types(
            defaults=(),
            overrides=(
                ("word/document.xml", WORD_DOCUMENT_CONTENT_TYPE),
                (part_name, "image/png"),
                ("word/media/private.png", "image/png"),
            ),
        ),
        "_rels/.rels": _relationship_xml(
            (
                (
                    "rIdDocument",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/officeDocument",
                    "word/document.xml",
                    None,
                ),
            )
        ),
        "word/document.xml": document_xml,
        "word/_rels/document.xml.rels": _relationship_xml(
            (
                (
                    "rIdLiveImage",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/image",
                    target,
                    None,
                ),
                (
                    "rIdPrivateImage",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/image",
                    "/word/media/private.png",
                    None,
                ),
            )
        ),
        part_name: b"live image",
        "word/media/private.png": b"CONFIDENTIAL PRIVATE IMAGE",
    }

    sanitized = _sanitize_members(members)

    relationships = sanitized["word/_rels/document.xml.rels"]
    assert _relationship_ids(relationships) == {"rIdLiveImage"}
    assert _relationship_targets(relationships) == {target}
    assert part_name in sanitized
    assert "word/media/private.png" not in sanitized
    assert part_name.encode() in sanitized["[Content_Types].xml"]
    assert b"private.png" not in sanitized["[Content_Types].xml"]


def test_sanitizer_does_not_parse_relationship_owner_binary() -> None:
    document_xml = f"""<w:document
 xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
 xmlns:r="{OFFICE_RELATIONSHIP_NAMESPACE}">
 <w:object r:id="rIdBinary"/>
 </w:document>""".encode()
    members = {
        "[Content_Types].xml": _typed_content_types(
            defaults=(),
            overrides=(
                ("word/document.xml", "application/xml"),
                ("word/embeddings/live.bin", "application/octet-stream"),
                ("word/media/private.png", "image/png"),
            ),
        ),
        "_rels/.rels": _relationship_xml(
            (
                (
                    "rIdDocument",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/officeDocument",
                    "word/document.xml",
                    None,
                ),
            )
        ),
        "word/document.xml": document_xml,
        "word/_rels/document.xml.rels": _relationship_xml(
            (
                (
                    "rIdBinary",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/oleObject",
                    "embeddings/live.bin",
                    None,
                ),
            )
        ),
        "word/embeddings/live.bin": (
            b"not XML: r:id='rIdPrivateImage'\x00\xff"
        ),
        "word/embeddings/_rels/live.bin.rels": _relationship_xml(
            (
                (
                    "rIdPrivateImage",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/image",
                    "../media/private.png",
                    None,
                ),
            )
        ),
        "word/media/private.png": b"CONFIDENTIAL BINARY DESCENDANT",
    }

    sanitized = _sanitize_members(members)

    assert "word/embeddings/live.bin" in sanitized
    assert (
        _relationship_targets(sanitized["word/embeddings/_rels/live.bin.rels"])
        == set()
    )
    assert "word/media/private.png" not in sanitized
    assert b"CONFIDENTIAL" not in b"".join(sanitized.values())


def test_sanitizer_fails_closed_for_unparseable_vml_owner() -> None:
    document_xml = f"""<w:document
 xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
 xmlns:r="{OFFICE_RELATIONSHIP_NAMESPACE}">
 <w:object r:id="rIdBrokenVml"/>
 </w:document>""".encode()
    members = {
        "[Content_Types].xml": _typed_content_types(
            defaults=(("vml", VML_CONTENT_TYPE),),
            overrides=(
                ("word/document.xml", "application/xml"),
                ("word/media/private.png", "image/png"),
            ),
        ),
        "_rels/.rels": _relationship_xml(
            (
                (
                    "rIdDocument",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/officeDocument",
                    "word/document.xml",
                    None,
                ),
            )
        ),
        "word/document.xml": document_xml,
        "word/_rels/document.xml.rels": _relationship_xml(
            (
                (
                    "rIdBrokenVml",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/vmlDrawing",
                    "drawings/broken.vml",
                    None,
                ),
            )
        ),
        "word/drawings/broken.vml": b"<v:shape",
        "word/drawings/_rels/broken.vml.rels": _relationship_xml(
            (
                (
                    "rIdPrivateImage",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/image",
                    "../media/private.png",
                    None,
                ),
            )
        ),
        "word/media/private.png": b"CONFIDENTIAL MALFORMED VML TARGET",
    }

    sanitized = _sanitize_members(members)

    assert "word/drawings/broken.vml" in sanitized
    assert (
        _relationship_targets(sanitized["word/drawings/_rels/broken.vml.rels"])
        == set()
    )
    assert "word/media/private.png" not in sanitized
    assert b"CONFIDENTIAL" not in b"".join(sanitized.values())


def test_sanitizer_keeps_live_embed_and_link_relationships() -> None:
    document_xml = f"""<w:document
 xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
 xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
 xmlns:r="{OFFICE_RELATIONSHIP_NAMESPACE}">
 <a:blip r:embed="rIdEmbeddedImage" r:link="rIdLinkedImage"/>
 </w:document>""".encode()
    members = {
        "[Content_Types].xml": _content_types("word/document.xml"),
        "_rels/.rels": _relationship_xml(
            (
                (
                    "rIdDocument",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/officeDocument",
                    "word/document.xml",
                    None,
                ),
            )
        ),
        "word/document.xml": document_xml,
        "word/_rels/document.xml.rels": _relationship_xml(
            (
                (
                    "rIdEmbeddedImage",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/image",
                    "media/live.png",
                    None,
                ),
                (
                    "rIdLinkedImage",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/image",
                    "https://images.example.test/live.png",
                    "External",
                ),
                (
                    "rIdUnusedImage",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/image",
                    "media/private.png",
                    None,
                ),
            )
        ),
        "word/media/live.png": b"live image",
        "word/media/private.png": b"private image",
    }

    sanitized = _sanitize_members(members)

    targets = _relationship_targets(sanitized["word/_rels/document.xml.rels"])
    assert targets == {
        "media/live.png",
        "https://images.example.test/live.png",
    }
    assert "word/media/live.png" in sanitized
    assert "word/media/private.png" not in sanitized


def test_sanitizer_keeps_only_live_picture_bullet_resources() -> None:
    word_namespace = (
        "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    )
    relationships = f'xmlns:r="{OFFICE_RELATIONSHIP_NAMESPACE}"'
    word = f'xmlns:w="{word_namespace}"'
    compatibility = (
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:w14="urn:test:w14"'
    )
    document_xml = (
        f"<w:document {word}><w:body><w:p><w:pPr>"
        '<w:pStyle w:val="ListBullet"/></w:pPr></w:p>'
        '<w:p><w:pPr><w:numPr><w:numId w:val="2"/></w:numPr>'
        "</w:pPr></w:p>"
        "</w:body></w:document>"
    ).encode()
    styles_xml = (
        f"<w:styles {word}>"
        '<w:style w:type="paragraph" w:styleId="Normal"/>'
        '<w:style w:type="paragraph" w:styleId="ListBullet">'
        '<w:basedOn w:val="Normal"/><w:pPr><w:numPr>'
        '<w:numId w:val="1"/></w:numPr></w:pPr></w:style>'
        "</w:styles>"
    ).encode()
    numbering_xml = (
        f"<w:numbering {word} {relationships} {compatibility}>"
        '<w:numPicBullet w:numPicBulletId="7"><w:pict '
        'r:id="rIdLivePicture"/></w:numPicBullet>'
        '<w:numPicBullet w:numPicBulletId="8"><w:pict '
        'r:id="rIdOverridePicture"/></w:numPicBullet>'
        '<w:numPicBullet w:numPicBulletId="9"><w:pict '
        'r:id="rIdPrivatePicture"/></w:numPicBullet>'
        '<w:numPicBullet w:numPicBulletId="10"><w:pict '
        'r:id="rIdCompatibilityPicture"/></w:numPicBullet>'
        '<w:abstractNum w:abstractNumId="10"><w:lvl w:ilvl="0">'
        '<w:lvlPicBulletId w:val="7"/></w:lvl><w:lvl w:ilvl="1">'
        '<w:lvlPicBulletId w:val="9"/></w:lvl><w:lvl w:ilvl="2">'
        '<mc:AlternateContent><mc:Choice Requires="w14">'
        '<w:lvlPicBulletId w:val="10"/></mc:Choice>'
        '<mc:Fallback><w:numFmt w:val="bullet"/></mc:Fallback>'
        "</mc:AlternateContent></w:lvl></w:abstractNum>"
        '<w:abstractNum w:abstractNumId="11"><w:lvl w:ilvl="0">'
        '<w:lvlPicBulletId w:val="9"/></w:lvl></w:abstractNum>'
        '<w:num w:numId="1"><w:abstractNumId w:val="10"/>'
        '<w:lvlOverride w:ilvl="0"><w:lvl w:ilvl="0">'
        '<w:lvlPicBulletId w:val="8"/></w:lvl></w:lvlOverride>'
        '<w:lvlOverride w:ilvl="1"><w:lvl w:ilvl="1">'
        '<w:numFmt w:val="bullet"/></w:lvl></w:lvlOverride></w:num>'
        '<w:num w:numId="2"><w:abstractNumId w:val="10"/>'
        '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="3"/>'
        "</w:lvlOverride>"
        '<w:lvlOverride w:ilvl="1"><w:lvl w:ilvl="1">'
        '<w:numFmt w:val="bullet"/></w:lvl></w:lvlOverride></w:num>'
        '<w:num w:numId="3"><w:abstractNumId w:val="11"/></w:num>'
        "</w:numbering>"
    ).encode()
    members = {
        "[Content_Types].xml": _typed_content_types(
            defaults=(),
            overrides=(
                ("word/document.xml", WORD_DOCUMENT_CONTENT_TYPE),
                ("word/styles.xml", "application/test-styles+xml"),
                ("word/numbering.xml", "application/test-numbering+xml"),
                ("word/media/live.png", "image/png"),
                ("word/media/override.png", "image/png"),
                ("word/media/private.png", "image/png"),
                ("word/media/compatibility.png", "image/png"),
            ),
        ),
        "_rels/.rels": _relationship_xml(
            (
                (
                    "rIdDocument",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/officeDocument",
                    "word/document.xml",
                    None,
                ),
            )
        ),
        "word/document.xml": document_xml,
        "word/_rels/document.xml.rels": _relationship_xml(
            (
                (
                    "rIdStyles",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/styles",
                    "styles.xml",
                    None,
                ),
                (
                    "rIdNumbering",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/numbering",
                    "numbering.xml",
                    None,
                ),
            )
        ),
        "word/styles.xml": styles_xml,
        "word/numbering.xml": numbering_xml,
        "word/_rels/numbering.xml.rels": _relationship_xml(
            (
                (
                    "rIdLivePicture",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/image",
                    "media/live.png",
                    None,
                ),
                (
                    "rIdOverridePicture",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/image",
                    "media/override.png",
                    None,
                ),
                (
                    "rIdPrivatePicture",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/image",
                    "media/private.png",
                    None,
                ),
                (
                    "rIdCompatibilityPicture",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/image",
                    "media/compatibility.png",
                    None,
                ),
            )
        ),
        "word/media/live.png": b"live picture bullet",
        "word/media/override.png": b"override picture bullet",
        "word/media/private.png": b"CONFIDENTIAL PRIVATE PICTURE BULLET",
        "word/media/compatibility.png": b"compatibility picture bullet",
    }

    sanitized = _sanitize_members(members)

    numbering = ElementTree.fromstring(sanitized["word/numbering.xml"])
    assert _relationship_ids(sanitized["word/_rels/numbering.xml.rels"]) == {
        "rIdLivePicture",
        "rIdOverridePicture",
        "rIdCompatibilityPicture",
    }
    assert {
        element.get(f"{{{word_namespace}}}numPicBulletId")
        for element in numbering.findall(f"{{{word_namespace}}}numPicBullet")
    } == {"7", "8", "10"}
    assert {
        element.get(f"{{{word_namespace}}}abstractNumId")
        for element in numbering.findall(f"{{{word_namespace}}}abstractNum")
    } == {"10"}
    assert {
        element.get(f"{{{word_namespace}}}numId")
        for element in numbering.findall(f"{{{word_namespace}}}num")
    } == {"1", "2"}
    assert "word/media/live.png" in sanitized
    assert "word/media/override.png" in sanitized
    assert "word/media/compatibility.png" in sanitized
    assert "word/media/private.png" not in sanitized
    assert b"private.png" not in sanitized["[Content_Types].xml"]
    assert b"CONFIDENTIAL" not in b"".join(sanitized.values())


def test_sanitizer_handles_missing_numbering_references() -> None:
    word_namespace = (
        "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    )
    word = f'xmlns:w="{word_namespace}"'
    relationships = f'xmlns:r="{OFFICE_RELATIONSHIP_NAMESPACE}"'
    document_xml = (
        f"<w:document {word}><w:body>"
        '<w:p><w:pPr><w:numPr><w:numId w:val="3"/></w:numPr>'
        "</w:pPr></w:p>"
        '<w:p><w:pPr><w:numPr><w:numId w:val="4"/></w:numPr>'
        "</w:pPr></w:p></w:body></w:document>"
    ).encode()
    numbering_xml = (
        f"<w:numbering {word} {relationships}>"
        '<w:numPicBullet w:numPicBulletId="30"><w:pict '
        'r:id="rIdUnrelatedPicture"/></w:numPicBullet>'
        '<w:abstractNum w:abstractNumId="20"><w:lvl w:ilvl="0">'
        '<w:lvlPicBulletId w:val="404"/></w:lvl></w:abstractNum>'
        '<w:num w:numId="3"><w:abstractNumId w:val="999"/></w:num>'
        '<w:num w:numId="4"><w:abstractNumId w:val="20"/>'
        '<w:lvlOverride w:ilvl="9"><w:startOverride w:val="2"/>'
        "</w:lvlOverride></w:num></w:numbering>"
    ).encode()
    members = {
        "[Content_Types].xml": _typed_content_types(
            defaults=(),
            overrides=(
                ("word/document.xml", WORD_DOCUMENT_CONTENT_TYPE),
                ("word/styles.xml", "application/test-styles+xml"),
                ("word/numbering.xml", "application/test-numbering+xml"),
                ("word/media/unrelated.png", "image/png"),
            ),
        ),
        "_rels/.rels": _relationship_xml(
            (
                (
                    "rIdDocument",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/officeDocument",
                    "word/document.xml",
                    None,
                ),
            )
        ),
        "word/document.xml": document_xml,
        "word/_rels/document.xml.rels": _relationship_xml(
            (
                (
                    "rIdStyles",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/styles",
                    "styles.xml",
                    None,
                ),
                (
                    "rIdNumbering",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/numbering",
                    "numbering.xml",
                    None,
                ),
            )
        ),
        "word/styles.xml": f"<w:styles {word}/>".encode(),
        "word/numbering.xml": numbering_xml,
        "word/_rels/numbering.xml.rels": _relationship_xml(
            (
                (
                    "rIdUnrelatedPicture",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/image",
                    "media/unrelated.png",
                    None,
                ),
            )
        ),
        "word/media/unrelated.png": b"unrelated picture bullet",
    }

    sanitized = _sanitize_members(members)

    numbering = ElementTree.fromstring(sanitized["word/numbering.xml"])
    assert numbering.findall(f"{{{word_namespace}}}numPicBullet") == []
    assert {
        element.get(f"{{{word_namespace}}}abstractNumId")
        for element in numbering.findall(f"{{{word_namespace}}}abstractNum")
    } == {"20"}
    assert {
        element.get(f"{{{word_namespace}}}numId")
        for element in numbering.findall(f"{{{word_namespace}}}num")
    } == {"3", "4"}
    assert (
        _relationship_ids(sanitized["word/_rels/numbering.xml.rels"]) == set()
    )
    assert "word/media/unrelated.png" not in sanitized
    assert b"unrelated.png" not in sanitized["[Content_Types].xml"]


def test_sanitizer_prunes_unreferenced_chart_relationship_chains() -> None:
    document_xml = f"""<w:document
 xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
 xmlns:r="{OFFICE_RELATIONSHIP_NAMESPACE}"><w:headerReference
 r:id="rIdHeader"/></w:document>""".encode()
    chart_xml = f"""<c:chartSpace
 xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"
 xmlns:r="{OFFICE_RELATIONSHIP_NAMESPACE}"><c:externalData
 r:id="rIdWorkbook"/></c:chartSpace>""".encode()
    members = {
        "[Content_Types].xml": _content_types(
            "word/document.xml",
            "word/header1.xml",
            "word/charts/private-chart.xml",
            "word/embeddings/private.xlsx",
        ),
        "_rels/.rels": _relationship_xml(
            (
                (
                    "rIdDocument",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/officeDocument",
                    "word/document.xml",
                    None,
                ),
            )
        ),
        "word/document.xml": document_xml,
        "word/_rels/document.xml.rels": _relationship_xml(
            (
                (
                    "rIdHeader",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/header",
                    "header1.xml",
                    None,
                ),
            )
        ),
        "word/header1.xml": b"<header/>",
        "word/_rels/header1.xml.rels": _relationship_xml(
            (
                (
                    "rIdPrivateChart",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/chart",
                    "charts/private-chart.xml",
                    None,
                ),
            )
        ),
        "word/charts/private-chart.xml": chart_xml,
        "word/charts/_rels/private-chart.xml.rels": _relationship_xml(
            (
                (
                    "rIdWorkbook",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/package",
                    "../embeddings/private.xlsx",
                    None,
                ),
            )
        ),
        "word/embeddings/private.xlsx": b"confidential chart workbook",
    }

    sanitized = _sanitize_members(members)

    assert "word/header1.xml" in sanitized
    assert "word/_rels/header1.xml.rels" in sanitized
    assert (
        _relationship_targets(sanitized["word/_rels/header1.xml.rels"]) == set()
    )
    for private_part in (
        "word/charts/private-chart.xml",
        "word/charts/_rels/private-chart.xml.rels",
        "word/embeddings/private.xlsx",
    ):
        assert private_part not in sanitized
    assert b"confidential chart workbook" not in b"".join(sanitized.values())
    assert b"private-chart.xml" not in sanitized["[Content_Types].xml"]
    assert b"private.xlsx" not in sanitized["[Content_Types].xml"]


def test_sanitizer_resolves_top_level_relationship_part_owners() -> None:
    document_xml = f"""<w:document
 xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
 xmlns:r="{OFFICE_RELATIONSHIP_NAMESPACE}"><w:customXml
 r:id="rIdLiveTop"/></w:document>""".encode()
    live_xml = f"""<live xmlns:r="{OFFICE_RELATIONSHIP_NAMESPACE}">
 <child r:id="rIdLiveChild"/></live>""".encode()
    members = {
        "[Content_Types].xml": _content_types(
            "word/document.xml",
            "live.xml",
            "live-child.bin",
            "private.xml",
            "private-child.bin",
        ),
        "_rels/.rels": _relationship_xml(
            (
                (
                    "rIdDocument",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/officeDocument",
                    "word/document.xml",
                    None,
                ),
            )
        ),
        "word/document.xml": document_xml,
        "word/_rels/document.xml.rels": _relationship_xml(
            (
                (
                    "rIdLiveTop",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/customXml",
                    "../live.xml",
                    None,
                ),
            )
        ),
        "live.xml": live_xml,
        "_rels/live.xml.rels": _relationship_xml(
            (
                (
                    "rIdLiveChild",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/image",
                    "live-child.bin",
                    None,
                ),
            )
        ),
        "live-child.bin": b"live bytes",
        "private.xml": b"<private/>",
        "_rels/private.xml.rels": _relationship_xml(
            (
                (
                    "rIdPrivateChild",
                    f"{OFFICE_RELATIONSHIP_NAMESPACE}/image",
                    "private-child.bin",
                    None,
                ),
            )
        ),
        "private-child.bin": b"confidential top-level bytes",
    }

    sanitized = _sanitize_members(members)

    assert "live.xml" in sanitized
    assert "_rels/live.xml.rels" in sanitized
    assert "live-child.bin" in sanitized
    for private_part in (
        "private.xml",
        "_rels/private.xml.rels",
        "private-child.bin",
    ):
        assert private_part not in sanitized
    assert b"confidential top-level bytes" not in b"".join(sanitized.values())
    assert b"private.xml" not in sanitized["[Content_Types].xml"]
    assert b"private-child.bin" not in sanitized["[Content_Types].xml"]


def test_writer_applies_layout_content_links_and_safe_metadata(
    tmp_path: Path,
) -> None:
    template_path = tmp_path / "template.docx"
    template(template_path)
    store = ArtefactStore(tmp_path / "artefacts")
    writer = WordWriter(store, available_fonts={"Aptos", "Aptos Display"})

    candidate = writer.write(
        content(),
        "Senior Python Engineer",
        "Acme/cv-generations/generation-1",
        template_path,
        layout(),
    )

    assert candidate.filename == ("Avery Morgan - Senior Python Engineer.docx")
    assert len(candidate.sha256) == 64
    document = Document(str(store.resolve(candidate.relative_path)))
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert texts == [
        "Avery Morgan",
        ("Software Engineer | avery@example.test | https://example.test/avery"),
        "Profile",
        "Builds typed Python services.",
        "Skills",
        "Python, FastAPI",
    ]
    assert "Private placeholder" not in " ".join(texts)
    assert all(
        "Private header" not in paragraph.text
        for paragraph in document.sections[0].header.paragraphs
    )
    section = document.sections[0]
    assert section.page_width == pytest.approx(Mm(210), abs=1_000)
    assert section.page_height == pytest.approx(Mm(297), abs=1_000)
    assert section.left_margin == pytest.approx(Inches(0.55), abs=1_000)
    assert document.styles["Normal"].font.name == "Aptos"
    assert document.styles["Heading 1"].font.name == "Aptos Display"
    bullet_style = document.paragraphs[-1].style
    assert bullet_style is not None
    assert bullet_style.name == "List Bullet"
    relationships = document.part.rels.values()
    targets = {
        relationship.target_ref
        for relationship in relationships
        if relationship.reltype.endswith("/hyperlink")
    }
    assert targets == {
        "mailto:avery@example.test",
        "https://example.test/avery",
    }
    properties = document.core_properties
    assert properties.title == ("Avery Morgan - Senior Python Engineer")
    assert properties.subject == "Tailored curriculum vitae"
    assert properties.author == "Avery Morgan"
    assert properties.last_modified_by == "JobHunter"
    assert properties.comments == (
        "Generated by JobHunter from validated CV content."
    )
    assert properties.category == "CV"
    assert properties.language == "en"
    assert properties.version == "1"
    assert properties.revision == 1
    assert properties.created is not None
    assert properties.created.year == 2000
    assert properties.modified is not None
    assert properties.modified.year == 2000


@pytest.mark.parametrize(
    ("name", "role", "expected"),
    [
        (
            "Avery Morgan",
            "Platform/Infrastructure Engineer",
            "Avery Morgan - Platform Infrastructure Engineer.docx",
        ),
        ("../", "../../", "Candidate - Role.docx"),
        ("CON", "Engineer", "_CON - Engineer.docx"),
    ],
)
def test_candidate_filename_is_portable(
    name: str, role: str, expected: str
) -> None:
    assert candidate_filename(name, role) == expected


@pytest.mark.parametrize(
    ("name", "role"),
    (
        ("é" * 100, "工程师" * 100),
        ("𐐀" * 60, "Engineer"),
    ),
)
def test_candidate_filename_reserves_atomic_utf8_headroom(
    name: str, role: str
) -> None:
    filename = candidate_filename(name, role)
    temporary_name = f"{filename}.tmp-{'0' * 16}"

    assert len(filename) <= 80
    assert len(filename.encode("utf-8")) <= MAX_CANDIDATE_BYTES
    assert len(temporary_name.encode("utf-8")) <= 255
    stem = filename.removesuffix(".docx")
    name_part, separator, role_part = stem.rpartition(" - ")
    assert separator == " - "
    assert name_part
    assert role_part
    assert filename.encode("utf-8").decode("utf-8") == filename


def test_candidate_filename_truncates_on_unicode_boundaries_and_keeps_role() -> (
    None
):
    filename = candidate_filename("𐐀" * 60, "Engineer")

    assert filename == f"{'𐐀' * 56} - En.docx"
    assert len(filename.encode("utf-8")) == MAX_CANDIDATE_BYTES


def test_writer_round_trips_overlong_multibyte_candidate_name(
    tmp_path: Path,
) -> None:
    template_path = tmp_path / "template.docx"
    template(template_path)
    store = ArtefactStore(tmp_path / "artefacts")
    original = content()
    long_name = original.identity.full_name.model_copy(
        update={"text": "𐐀" * 60}
    )
    long_content = original.model_copy(
        update={
            "identity": original.identity.model_copy(
                update={"full_name": long_name}
            )
        }
    )

    candidate = WordWriter(store).write(
        long_content,
        "Engineer",
        "safe",
        template_path,
        layout(),
    )

    assert len(candidate.filename.encode("utf-8")) <= MAX_CANDIDATE_BYTES
    reopened = Document(BytesIO(store.read_bytes(candidate.relative_path)))
    assert reopened.paragraphs[0].text == "𐐀" * 60


def test_writer_rejects_missing_styles_fonts_and_unsafe_output(
    tmp_path: Path,
) -> None:
    template_path = tmp_path / "template.docx"
    template(template_path)
    document = Document(str(template_path))
    title = document.styles["Title"]
    title.element.getparent().remove(title.element)
    document.save(str(template_path))
    store = ArtefactStore(tmp_path / "artefacts")

    with pytest.raises(WordWriterError, match="missing required style"):
        WordWriter(store).write(
            content(), "Engineer", "safe", template_path, layout()
        )

    template(template_path)
    with pytest.raises(WordWriterError, match="fonts are unavailable"):
        WordWriter(store, available_fonts={"Aptos"}).write(
            content(), "Engineer", "safe", template_path, layout()
        )

    with pytest.raises(UnsafeArtefactPathError):
        WordWriter(store).write(
            content(), "Engineer", "../escape", template_path, layout()
        )
    assert list(store.root.rglob("*.docx")) == []


def test_writer_rejects_corrupt_template(tmp_path: Path) -> None:
    template_path = tmp_path / "template.docx"
    template_path.write_bytes(b"not a DOCX")

    with pytest.raises(WordWriterError, match="corrupt or unreadable"):
        WordWriter(ArtefactStore(tmp_path / "artefacts")).write(
            content(), "Engineer", "safe", template_path, layout()
        )


def test_writer_sets_east_asia_font_names(tmp_path: Path) -> None:
    template_path = tmp_path / "template.docx"
    template(template_path)
    store = ArtefactStore(tmp_path / "artefacts")

    candidate = WordWriter(store).write(
        content(), "Engineer", "safe", template_path, layout()
    )

    document = Document(str(store.resolve(candidate.relative_path)))
    normal_fonts = document.styles["Normal"].element.rPr.rFonts
    assert normal_fonts.get(qn("w:eastAsia")) == "Aptos"


def test_writer_removes_private_story_content_and_relationships(
    tmp_path: Path,
) -> None:
    template_path = tmp_path / "template.docx"
    private_template_content(template_path)
    with zipfile.ZipFile(template_path) as archive:
        private_template_package = {
            member.filename: archive.read(member.filename)
            for member in archive.infolist()
        }
    for private_part in (
        "docProps/custom.xml",
        "docProps/thumbnail.jpeg",
        "word/comments.xml",
        "word/footnotes.xml",
        "word/glossary/document.xml",
    ):
        assert private_part in private_template_package
    assert b"CONFIDENTIAL" in b"".join(private_template_package.values())
    store = ArtefactStore(tmp_path / "artefacts")

    candidate = WordWriter(store).write(
        content(), "Engineer", "safe", template_path, layout()
    )

    candidate_path = store.resolve(candidate.relative_path)
    document = Document(str(candidate_path))
    for section in document.sections:
        assert section.header.tables == []
        assert section.footer.tables == []
    with zipfile.ZipFile(candidate_path) as archive:
        package = {
            member.filename: archive.read(member.filename)
            for member in archive.infolist()
        }
    visible_xml = b"".join(
        data for name, data in package.items() if name.endswith(".xml")
    )
    relationships = b"".join(
        data for name, data in package.items() if name.endswith(".rels")
    )
    assert b"CONFIDENTIAL" not in b"".join(package.values())
    assert b"Private body placeholder" not in visible_xml
    assert b"Private header table" not in visible_xml
    assert b"Private footer table" not in visible_xml
    assert b"private.example.test" not in relationships
    assert b"private-template.example.test" not in relationships
    assert b"client-list.xlsx" not in relationships
    for private_relationship_id in (
        b"rIdPrivateTemplate",
        b"rIdPrivateDataSource",
        b"rIdChoiceTemplate",
        b"rIdChoiceDataSource",
        b"rIdFallbackTemplate",
        b"rIdFallbackDataSource",
    ):
        assert private_relationship_id not in relationships
    assert "word/media/private-image.png" not in package
    assert "word/embeddings/private-object.bin" not in package
    assert "word/activeX/private-control.xml" not in package
    assert "word/activeX/_rels/private-control.xml.rels" not in package
    assert "word/embeddings/private-active.bin" not in package
    assert "word/media/private-preview.png" not in package
    for private_part in (
        "docProps/custom.xml",
        "docProps/thumbnail.jpeg",
        "word/comments.xml",
        "word/_rels/comments.xml.rels",
        "word/footnotes.xml",
        "word/glossary/document.xml",
        "word/media/private-comment.png",
    ):
        assert private_part not in package
    for rendering_part in (
        "docProps/core.xml",
        "word/styles.xml",
        "word/settings.xml",
        "word/numbering.xml",
        "word/fontTable.xml",
        "word/theme/theme1.xml",
    ):
        assert rendering_part in package
    content_types = package["[Content_Types].xml"]
    assert b"private-control.xml" not in content_types
    assert b"private-active.bin" not in content_types
    for private_override in (
        b"custom.xml",
        b"thumbnail.jpeg",
        b"comments.xml",
        b"footnotes.xml",
        b"glossary/document.xml",
        b"private-comment.png",
    ):
        assert private_override not in content_types
    settings = package["word/settings.xml"].decode("utf-8")
    settings_root = ElementTree.fromstring(settings)
    private_settings_tags = {
        qn("w:attachedTemplate"),
        qn("w:docVars"),
        qn("w:mailMerge"),
    }
    assert all(
        element.tag not in private_settings_tags
        for element in settings_root.iter()
    )
    markup_compatibility_namespace = (
        "http://schemas.openxmlformats.org/markup-compatibility/2006"
    )
    assert (
        settings_root.find(f".//{{{markup_compatibility_namespace}}}Choice")
        is not None
    )
    assert (
        settings_root.find(f".//{{{markup_compatibility_namespace}}}Fallback")
        is not None
    )
    assert {
        zoom.get(qn("w:percent")) for zoom in settings_root.iter(qn("w:zoom"))
    }.issuperset({"85", "90"})
    assert any(
        setting.get(qn("w:name")) == "safe"
        and setting.get(qn("w:uri")) == "urn:safe"
        and setting.get(qn("w:val")) == "1"
        for setting in settings_root.iter(qn("w:compatSetting"))
    )
    ignorable = settings_root.get(
        f"{{{markup_compatibility_namespace}}}Ignorable"
    )
    assert ignorable is not None
    for prefix in ignorable.split():
        assert f"xmlns:{prefix}=" in settings
    core_properties = package["docProps/core.xml"]
    assert b"contentStatus" not in core_properties
    assert b"CONFIDENTIAL REVIEW STATUS" not in core_properties
    assert b"CONFIDENTIAL UNKNOWN CORE PROPERTY" not in core_properties
    core_root = ElementTree.fromstring(core_properties)
    assert all(
        child.tag != "{urn:jobhunter:test-private}metadata"
        for child in core_root
    )


@pytest.mark.parametrize("relationship_kind", ("internal", "external", "none"))
def test_writer_removes_private_document_backgrounds(
    tmp_path: Path,
    relationship_kind: Literal["external", "internal", "none"],
) -> None:
    template_path = tmp_path / "template.docx"
    background_template(template_path, relationship_kind)
    store = ArtefactStore(tmp_path / "artefacts")

    candidate = WordWriter(store).write(
        content(), "Engineer", "safe", template_path, layout()
    )

    candidate_path = store.resolve(candidate.relative_path)
    with zipfile.ZipFile(candidate_path) as archive:
        package = {
            member.filename: archive.read(member.filename)
            for member in archive.infolist()
        }
    document_root = ElementTree.fromstring(package["word/document.xml"])
    assert document_root.find(qn("w:background")) is None
    relationships = package["word/_rels/document.xml.rels"]
    assert "rIdPrivateBackground" not in _relationship_ids(relationships)
    assert b"private-background.example.test" not in relationships
    assert "word/media/private-background.png" not in package
    assert b"private-background.png" not in package["[Content_Types].xml"]
    assert b"PRIVATE BACKGROUND SENTINEL" not in b"".join(package.values())

    document = Document(str(candidate_path))
    assert document.paragraphs[0].text == "Avery Morgan"
    assert document.sections[0].page_width == pytest.approx(Mm(210), abs=1_000)
    assert document.styles["Normal"].font.name == "Aptos"


def test_writer_removes_private_section_metadata_and_printer_settings(
    tmp_path: Path,
) -> None:
    template_path = tmp_path / "template.docx"
    private_section_metadata_template(template_path)
    store = ArtefactStore(tmp_path / "artefacts")

    candidate = WordWriter(store).write(
        content(), "Engineer", "safe", template_path, layout()
    )

    candidate_path = store.resolve(candidate.relative_path)
    with zipfile.ZipFile(candidate_path) as archive:
        package = {
            member.filename: archive.read(member.filename)
            for member in archive.infolist()
        }
    document_xml = package["word/document.xml"]
    relationships = package["word/_rels/document.xml.rels"]
    content_types = package["[Content_Types].xml"]
    assert b"sectPrChange" not in document_xml
    assert b"PRIVATE-RSID" not in document_xml
    assert b"Private Reviewer" not in document_xml
    assert b"2026-08-09T12:00:00Z" not in document_xml
    assert b"printerSettings" not in document_xml
    assert b"rIdPrivatePrinterSettings" not in relationships
    assert b"printerSettings1.bin" not in content_types
    assert "word/printerSettings1.bin" not in package

    document = Document(str(candidate_path))
    section = document.sections[0]
    assert section.orientation is WD_ORIENT.PORTRAIT
    assert section.page_width == pytest.approx(Mm(210), abs=1_000)
    assert section.page_height == pytest.approx(Mm(297), abs=1_000)


def test_writer_bounds_and_xml_sanitizes_derived_metadata_title(
    tmp_path: Path,
) -> None:
    template_path = tmp_path / "template.docx"
    template(template_path)
    store = ArtefactStore(tmp_path / "artefacts")
    role = "Principal\x01 Platform Engineer " + ("X" * 400)

    candidate = WordWriter(store).write(
        content(), role, "safe", template_path, layout()
    )

    candidate_path = store.resolve(candidate.relative_path)
    document = Document(str(candidate_path))
    title = document.core_properties.title
    assert title is not None
    assert len(title) == 255
    assert title.startswith("Avery Morgan - Principal Platform Engineer ")
    assert "\x01" not in title
    with zipfile.ZipFile(candidate_path) as archive:
        core_properties = archive.read("docProps/core.xml")
    assert b"\x01" not in core_properties


def test_writer_emits_byte_deterministic_docx_archives(tmp_path: Path) -> None:
    template_path = tmp_path / "template.docx"
    template(template_path)
    store = ArtefactStore(tmp_path / "artefacts")
    writer = WordWriter(store)

    first = writer.write(
        content(), "Engineer", "first", template_path, layout()
    )
    second = writer.write(
        content(), "Engineer", "second", template_path, layout()
    )

    first_bytes = store.read_bytes(first.relative_path)
    second_bytes = store.read_bytes(second.relative_path)
    assert first_bytes == second_bytes
    assert first.sha256 == second.sha256
    assert hashlib.sha256(first_bytes).hexdigest() == first.sha256
    with zipfile.ZipFile(BytesIO(first_bytes)) as archive:
        members = archive.infolist()
        assert [member.filename for member in members] == sorted(
            member.filename for member in members
        )
        assert all(
            member.date_time == (2000, 1, 1, 0, 0, 0) for member in members
        )
        assert all(member.create_system == 3 for member in members)
        assert all(member.external_attr == 0o600 << 16 for member in members)
        assert all(member.extra == b"" for member in members)
        assert all(member.comment == b"" for member in members)

    reopened = Document(BytesIO(first_bytes))
    assert reopened.paragraphs[0].text == "Avery Morgan"


@pytest.mark.parametrize(
    ("size", "width", "height"),
    [
        ("A4", Mm(210), Mm(297)),
        ("Letter", Inches(8.5), Inches(11)),
    ],
)
def test_writer_forces_portrait_layout_for_every_template_orientation(
    tmp_path: Path,
    size: str,
    width: int,
    height: int,
) -> None:
    template_path = tmp_path / "template.docx"
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    document.save(str(template_path))
    configured_layout = layout().model_copy(
        update={"page": layout().page.model_copy(update={"size": size})}
    )
    store = ArtefactStore(tmp_path / "artefacts")

    candidate = WordWriter(store).write(
        content(), "Engineer", "safe", template_path, configured_layout
    )

    rendered_section = Document(
        str(store.resolve(candidate.relative_path))
    ).sections[0]
    assert rendered_section.orientation is WD_ORIENT.PORTRAIT
    assert rendered_section.page_width == pytest.approx(width, abs=1_000)
    assert rendered_section.page_height == pytest.approx(height, abs=1_000)


def test_writer_uses_configured_section_order_and_link_color(
    tmp_path: Path,
) -> None:
    template_path = tmp_path / "template.docx"
    template(template_path)
    store = ArtefactStore(tmp_path / "artefacts")
    configured = layout().model_copy(
        update={
            "section_order": (
                "skills",
                "summary",
                "experience",
                "projects",
                "education",
                "certifications",
                "additional",
            )
        }
    )

    candidate = WordWriter(store).write(
        content(), "Engineer", "safe", template_path, configured
    )

    document = Document(str(store.resolve(candidate.relative_path)))
    assert [paragraph.text for paragraph in document.paragraphs][2:] == [
        "Skills",
        "Python, FastAPI",
        "Profile",
        "Builds typed Python services.",
    ]
    hyperlink_colors = document._element.xpath(
        ".//w:hyperlink/w:r/w:rPr/w:color/@w:val"
    )
    assert hyperlink_colors == ["2563EB", "2563EB"]
